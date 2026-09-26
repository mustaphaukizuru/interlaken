"""
GET /api/v1/accounts/admin/students/export/ on the Data Ops export contract
(C3): the roster list with another renderer (CSV/XLSX/PDF), honouring ``q``,
the filters (Spanish aliases included), ``ordering`` and ``?ids=``; audited,
admin only, no medical data, bounded queries.
"""

import csv
import io
from decimal import Decimal

import pytest
from django.urls import reverse

from apps.accounts.factories import AdminFactory, ParentFactory, StudentProfileFactory
from apps.core.models import AuditLog

pytestmark = pytest.mark.django_db

URL = reverse("export-students")


def _body(resp) -> bytes:
    return b"".join(resp.streaming_content) if resp.streaming else resp.content


def _rows(resp):
    text = _body(resp).decode("utf-8")
    assert text.startswith("﻿")
    return list(csv.reader(io.StringIO(text.lstrip("﻿"))))


def _names(resp):
    return [r[0] for r in _rows(resp)[1:]]


@pytest.fixture
def roster():
    from apps.cafeteria.models import CafeteriaBalance

    a = StudentProfileFactory(
        user__first_name="Ana", user__last_name="Zuñiga", student_id="09931", grade="1° Primaria"
    )
    b = StudentProfileFactory(
        user__first_name="Beto",
        user__last_name="Alvarado",
        student_id="09932",
        grade="2° Secundaria",
        group="B",
        status="on_leave",
        is_active=False,
    )
    c = StudentProfileFactory(
        user__first_name="Carla", user__last_name="Mora", student_id="09933", grade="1° Primaria"
    )
    CafeteriaBalance.objects.create(student=a, balance=Decimal("120.50"))
    CafeteriaBalance.objects.create(student=c, balance=Decimal("15.00"))
    return a, b, c


class TestCsv:
    def test_headers_bom_guardians_and_balance(self, admin_client, roster):
        a, *_ = roster
        a.parents.add(ParentFactory(), ParentFactory())
        resp = admin_client.get(URL)
        assert resp.status_code == 200
        assert resp["Content-Type"].startswith("text/csv")
        assert "attachment" in resp["Content-Disposition"]
        rows = _rows(resp)
        assert rows[0][:6] == ["Alumno", "Matrícula", "Código Loyverse", "Grado", "Grupo", "Estado"]
        by_name = {r[0]: dict(zip(rows[0], r, strict=True)) for r in rows[1:]}
        assert by_name["Ana Zuñiga"]["Tutores"] == "2"
        assert by_name["Ana Zuñiga"]["Saldo cafetería"] == "120.50"
        assert by_name["Beto Alvarado"]["Saldo cafetería"] == "0.00"
        assert by_name["Beto Alvarado"]["Estado"] == "Baja temporal"

    def test_honours_q_filters_and_ordering(self, admin_client, roster):
        assert _names(admin_client.get(URL, {"q": "ci09932"})) == ["Beto Alvarado"]
        assert _names(admin_client.get(URL, {"search": "mora"})) == ["Carla Mora"]
        assert _names(admin_client.get(URL, {"nivel": "primaria", "ordering": "-balance"})) == [
            "Ana Zuñiga",
            "Carla Mora",
        ]
        assert _names(admin_client.get(URL, {"estado": "on_leave"})) == ["Beto Alvarado"]
        assert _names(admin_client.get(URL, {"status": "active", "ordering": "name"})) == [
            "Carla Mora",
            "Ana Zuñiga",
        ]

    def test_ids_exports_only_the_selection(self, admin_client, roster):
        a, _, c = roster
        resp = admin_client.get(URL, {"ids": f"{c.id},{a.id}"})
        assert sorted(_names(resp)) == ["Ana Zuñiga", "Carla Mora"]

    def test_export_is_audited_with_filters_and_count(self, admin_client, roster):
        admin_client.get(URL, {"fmt": "csv", "grado": "1° Primaria"})
        log = AuditLog.objects.get(object_type="export:students")
        assert log.action == "export" and log.object_id == "csv"
        assert log.changes["rows"] == 2 and log.changes["filters"]["grado"] == "1° Primaria"


class TestFormats:
    def test_xlsx_has_real_numbers(self, admin_client, roster):
        from openpyxl import load_workbook

        resp = admin_client.get(URL, {"fmt": "xlsx", "ordering": "-balance"})
        assert resp.status_code == 200
        ws = load_workbook(io.BytesIO(_body(resp))).worksheets[0]
        header = [c.value for c in ws[1]]
        col = header.index("Saldo cafetería")
        assert ws.cell(row=2, column=col + 1).value == 120.5

    def test_pdf(self, admin_client, roster):
        resp = admin_client.get(URL, {"fmt": "pdf"})
        assert resp.status_code == 200 and _body(resp).startswith(b"%PDF")

    def test_bad_format_is_400(self, admin_client):
        assert admin_client.get(URL, {"fmt": "doc"}).status_code == 400


class TestAccessAndBudget:
    def test_parent_is_403(self, api_client, parent_user):
        api_client.force_authenticate(user=parent_user)
        assert api_client.get(URL).status_code == 403

    def test_anonymous_is_401(self, api_client):
        assert api_client.get(URL).status_code == 401

    @pytest.mark.parametrize("n", [3, 9])
    def test_constant_queries(self, api_client, django_assert_max_num_queries, n):
        parent = ParentFactory()
        for _ in range(n):
            StudentProfileFactory(parents=[parent], user__password=None)
        api_client.force_authenticate(AdminFactory())
        # count + rows (user, Loyverse snapshot, wallet, guardian count) + audit insert
        with django_assert_max_num_queries(4):
            resp = api_client.get(URL, {"fmt": "csv"})
            body = _body(resp)
        assert body.count(b"\r\n") == n + 1
