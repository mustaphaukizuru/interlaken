import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.factories import AdminFactory, ParentFactory
from apps.admissions.models import PreRegistration, Registration
from apps.core.models import AuditLog

pytestmark = pytest.mark.django_db


def _pre(**o):
    base = dict(child_first_name='Ana', child_last_name='Pérez', child_dob='2020-01-01', level='preescolar',
                grade_applying='Preescolar 1°', parent_name='Luis', parent_email='l@test.mx', parent_phone='55')
    base.update(o)
    return PreRegistration.objects.create(**base)


def _reg(**o):
    base = dict(child_first_name='Beto', child_last_name='Ruiz', child_dob='2018-01-01', level='primaria',
                grade_applying='Primaria 1°', parent1_name='Marta', parent1_email='m@test.mx',
                parent1_phone='55', status=Registration.Status.SUBMITTED, child_curp='RUXB180101HDFZZZ01',
                allergies='Nuez', blood_type='O+')
    base.update(o)
    return Registration.objects.create(**base)


def _csv(resp):
    return b''.join(resp.streaming_content).decode('utf-8-sig')


def test_preregistration_export_csv_honours_filters_and_is_audited(api_client):
    _pre()
    _pre(status='contacted', child_first_name='Beto')
    api_client.force_authenticate(user=AdminFactory())
    resp = api_client.get(reverse('pre-register-export'), {'status': 'contacted'})
    assert resp.status_code == 200 and resp['Content-Type'].startswith('text/csv')
    body = _csv(resp)
    assert 'Beto' in body and 'Ana Pérez' not in body
    log = AuditLog.objects.get(object_type='export:pre-registros')
    assert log.object_id == 'csv' and log.changes['rows'] == 1
    assert log.changes['filters'] == {'status': 'contacted'}


def test_preregistration_export_q_level_ids_and_xlsx(api_client):
    a = _pre(child_first_name='Lucía', level='primaria')
    b = _pre(child_first_name='Pedro', level='primaria')
    _pre(child_first_name='Zoe', level='secundaria')
    api_client.force_authenticate(user=AdminFactory())
    body = _csv(api_client.get(reverse('pre-register-export'), {'q': 'lucía'}))
    assert 'Lucía' in body and 'Pedro' not in body
    body = _csv(api_client.get(reverse('pre-register-export'), {'level': 'primaria', 'ordering': 'child'}))
    assert 'Zoe' not in body and body.index('Lucía') < body.index('Pedro')
    body = _csv(api_client.get(reverse('pre-register-export'), {'ids': f'{b.pk}'}))
    assert 'Pedro' in body and 'Lucía' not in body
    resp = api_client.get(reverse('pre-register-export'), {'fmt': 'xlsx', 'ids': f'{a.pk},{b.pk}'})
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws.max_row == 3 and ws['A1'].value == 'Fecha'
    pdf = api_client.get(reverse('pre-register-export'), {'fmt': 'pdf'})
    assert pdf.status_code == 200 and pdf.content.startswith(b'%PDF')


def test_registration_export_has_no_medical_fields_and_honours_q(api_client):
    _reg()
    _reg(child_first_name='Carla', child_curp='', parent1_email='c@test.mx')
    api_client.force_authenticate(user=AdminFactory())
    resp = api_client.get(reverse('register-export'), {'q': 'RUXB18'})
    body = _csv(resp)
    assert 'Beto Ruiz' in body and 'Carla' not in body
    assert 'Nuez' not in body and 'O+' not in body
    header = body.splitlines()[0]
    for word in ('Alergias', 'Sangre', 'Estatura', 'Peso'):
        assert word not in header
    assert AuditLog.objects.filter(object_type='export:inscripciones').exists()


def test_export_cannot_create_a_preregistration(api_client):
    resp = api_client.post(reverse('pre-register-export'), {}, format='json')
    assert resp.status_code in (401, 403, 405)
    assert PreRegistration.objects.count() == 0


def test_exports_forbidden_for_families(api_client):
    api_client.force_authenticate(user=ParentFactory())
    for name in ('pre-register-export', 'register-export', 'core-admin-audit-export'):
        assert api_client.get(reverse(name)).status_code == 403


def test_audit_export(api_client):
    api_client.force_authenticate(user=AdminFactory())
    _pre()  # creates audit rows via signals
    resp = api_client.get(reverse('core-admin-audit-export'))
    assert resp.status_code == 200 and resp.content.decode('utf-8-sig').startswith('Fecha,')
