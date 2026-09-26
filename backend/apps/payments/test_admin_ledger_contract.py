"""
Data Ops Phase 7: the admin Pagos ledger on the list contract and its export.

GET /api/v1/payments/admin/          q, status/gateway/type/student/dates, ordering, page_size
GET /api/v1/payments/admin/export/   same view as CSV / XLSX / PDF, ?ids=, audited, capped
"""

import io
from datetime import datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

import pytest
from django.urls import reverse

from apps.accounts.factories import AdminFactory, ParentFactory, StudentProfileFactory
from apps.cafeteria.models import TopUpRequest
from apps.core.models import AuditLog
from apps.payments.filters import ADMIN_PAYMENT_ORDERING
from apps.payments.models import Payment

pytestmark = pytest.mark.django_db

LIST = reverse("admin-payments")
EXPORT = reverse("admin-payments-export")
MX = ZoneInfo("America/Mexico_City")

# Mirrors PAYMENTS_ORDERING_KEYS in frontend/src/hooks/queries/payments.ts.
FRONTEND_ORDERING_KEYS = {"date", "amount", "status", "gateway", "student"}


def _pay(student, amount, status, *, gateway=Payment.Gateway.GLOBAL_PAYMENTS, ref="", when=None):
    topup = TopUpRequest.objects.create(
        student=student, amount=Decimal(amount), method=TopUpRequest.Method.ONLINE
    )
    payer = student.parents.first() or ParentFactory()
    p = Payment.objects.create(
        user=payer,
        payment_type=Payment.Type.CAFETERIA,
        amount=Decimal(amount),
        related_topup=topup,
        status=status,
        gateway=gateway,
        gateway_ref=ref,
    )
    if when is not None:
        Payment.objects.filter(pk=p.pk).update(created_at=when)
    return p


@pytest.fixture
def ledger():
    ana = StudentProfileFactory(
        user__first_name="Ana", user__last_name="Zuñiga", student_id="09931"
    )
    beto = StudentProfileFactory(
        user__first_name="Beto", user__last_name="Alvarado", student_id="09932"
    )
    return {
        "p1": _pay(
            ana,
            "200.00",
            Payment.Status.SUCCESS,
            ref="GP-111",
            when=datetime(2026, 3, 1, 10, tzinfo=MX),
        ),
        "p2": _pay(
            ana,
            "50.00",
            Payment.Status.FAILED,
            gateway=Payment.Gateway.BANORTE,
            ref="BN-222",
            when=datetime(2026, 3, 15, 10, tzinfo=MX),
        ),
        "p3": _pay(
            beto,
            "900.00",
            Payment.Status.SUCCESS,
            ref="GP-333",
            when=datetime(2026, 4, 2, 10, tzinfo=MX),
        ),
        "ana": ana,
        "beto": beto,
    }


def _ids(resp):
    return [r["id"] for r in resp.data["results"]]


def test_ordering_keys_match_the_frontend_constant():
    assert set(ADMIN_PAYMENT_ORDERING) == FRONTEND_ORDERING_KEYS


class TestList:
    def test_parent_is_403(self, api_client):
        api_client.force_authenticate(ParentFactory())
        assert api_client.get(LIST).status_code == 403
        assert api_client.get(EXPORT).status_code == 403

    def test_default_is_newest_first(self, admin_client, ledger):
        assert _ids(admin_client.get(LIST)) == [ledger["p3"].id, ledger["p2"].id, ledger["p1"].id]

    def test_q_over_student_matricula_and_gateway_refs(self, admin_client, ledger):
        assert _ids(admin_client.get(LIST, {"q": "alvarado"})) == [ledger["p3"].id]
        assert set(_ids(admin_client.get(LIST, {"q": "09931"}))) == {
            ledger["p1"].id,
            ledger["p2"].id,
        }
        assert _ids(admin_client.get(LIST, {"q": "ci09932"})) == [ledger["p3"].id]
        assert _ids(admin_client.get(LIST, {"q": "BN-222"})) == [ledger["p2"].id]

    def test_filters(self, admin_client, ledger):
        assert _ids(admin_client.get(LIST, {"status": "failed"})) == [ledger["p2"].id]
        assert _ids(admin_client.get(LIST, {"gateway": "banorte"})) == [ledger["p2"].id]
        assert _ids(admin_client.get(LIST, {"student": ledger["beto"].id})) == [ledger["p3"].id]
        assert _ids(admin_client.get(LIST, {"from": "2026-03-10", "to": "2026-03-31"})) == [
            ledger["p2"].id
        ]
        assert len(_ids(admin_client.get(LIST, {"type": "cafeteria"}))) == 3

    def test_unknown_filter_values_are_ignored_not_400(self, admin_client, ledger):
        for params in (
            {"status": "bogus"},
            {"gateway": "sandbox"},
            {"student": "x"},
            {"from": "2026-02-30"},
        ):
            resp = admin_client.get(LIST, params)
            assert resp.status_code == 200 and resp.data["count"] == 3, params

    def test_ordering(self, admin_client, ledger):
        assert _ids(admin_client.get(LIST, {"ordering": "amount"})) == [
            ledger["p2"].id,
            ledger["p1"].id,
            ledger["p3"].id,
        ]
        # student: last name then first name (Alvarado before Zuñiga), -pk tiebreak.
        assert _ids(admin_client.get(LIST, {"ordering": "student"})) == [
            ledger["p3"].id,
            ledger["p2"].id,
            ledger["p1"].id,
        ]

    def test_page_size(self, admin_client, ledger):
        resp = admin_client.get(LIST, {"page_size": 2})
        assert resp.data["count"] == 3 and len(resp.data["results"]) == 2

    @pytest.mark.parametrize("n", [3, 9])
    def test_query_budget_is_constant(self, api_client, django_assert_num_queries, n):
        for _ in range(n):
            _pay(StudentProfileFactory(), "10.00", Payment.Status.SUCCESS)
        api_client.force_authenticate(AdminFactory())
        with django_assert_num_queries(2):  # count + page (payer, topup, student, user joined)
            resp = api_client.get(LIST, {"q": "a", "ordering": "-student", "page_size": 50})
        assert resp.status_code == 200


class TestExport:
    def test_csv_honours_filters_and_is_audited(self, admin_client, ledger):
        resp = admin_client.get(EXPORT, {"fmt": "csv", "status": "success", "ordering": "amount"})
        assert resp.status_code == 200
        assert resp["Content-Type"].startswith("text/csv")
        body = b"".join(resp.streaming_content).decode("utf-8-sig")
        lines = body.strip().split("\r\n")
        assert lines[0].startswith("Fecha,Alumno,Matrícula")
        assert len(lines) == 3 and "200.00" in lines[1] and "900.00" in lines[2]
        entry = AuditLog.objects.get(object_type="export:payments")
        assert entry.changes["rows"] == 2 and entry.changes["filters"]["status"] == "success"

    def test_selected_ids_only(self, admin_client, ledger):
        resp = admin_client.get(EXPORT, {"fmt": "csv", "ids": f"{ledger['p2'].id}"})
        body = b"".join(resp.streaming_content).decode("utf-8-sig")
        assert "BN-222" in body and "GP-111" not in body

    def test_xlsx_has_real_numbers(self, admin_client, ledger):
        from openpyxl import load_workbook

        resp = admin_client.get(EXPORT, {"fmt": "xlsx"})
        assert resp.status_code == 200
        ws = load_workbook(io.BytesIO(resp.content)).active
        header = [c.value for c in ws[1]]
        amount_col = header.index("Monto")
        amounts = sorted(float(row[amount_col].value) for row in ws.iter_rows(min_row=2))
        assert amounts == [50.0, 200.0, 900.0]

    def test_pdf(self, admin_client, ledger):
        resp = admin_client.get(EXPORT, {"fmt": "pdf"})
        assert resp.status_code == 200 and resp.content.startswith(b"%PDF")

    def test_over_cap_is_413(self, admin_client, ledger, monkeypatch):
        from apps.payments import filters

        monkeypatch.setattr(filters.PAYMENT_EXPORT_SPEC, "row_cap", 2)
        resp = admin_client.get(EXPORT, {"fmt": "csv"})
        assert resp.status_code == 413
        assert "límite es 2" in resp.data["detail"]
