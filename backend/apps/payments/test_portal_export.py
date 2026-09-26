"""Data Ops Phase 9: family Pagos list ordering and export.

``PaymentHistoryExportView`` subclasses ``PaymentHistoryView``: the file is
the list the family is looking at (scoping, ``status``/``student``/dates and
the whitelisted ``ordering``), capped at 10,000 rows, audited with
``record_export`` and throttled under ``portal-export``.
"""

import io
from decimal import Decimal

import pytest
from django.urls import reverse

from apps.accounts.factories import ParentFactory, StudentProfileFactory
from apps.cafeteria.models import TopUpRequest
from apps.core.models import AuditLog
from apps.payments import views as payment_views
from apps.payments.models import Payment
from apps.payments.portal_exports import FAMILY_PAYMENTS_EXPORT_SPEC

pytestmark = pytest.mark.django_db


def _pay(student, payer, amount, status, ref):
    topup = TopUpRequest.objects.create(
        student=student, amount=Decimal(amount), method=TopUpRequest.Method.ONLINE
    )
    return Payment.objects.create(
        user=payer,
        payment_type=Payment.Type.CAFETERIA,
        amount=Decimal(amount),
        related_topup=topup,
        status=status,
        gateway_tx_id=ref,
    )


def _csv_rows(resp) -> list[str]:
    body = b"".join(resp.streaming_content).decode("utf-8-sig")
    return [line for line in body.split("\r\n")[1:] if line]


@pytest.fixture
def two_families(api_client):
    parent_a = ParentFactory()
    parent_b = ParentFactory()
    a = StudentProfileFactory(parents=[parent_a])
    b = StudentProfileFactory(parents=[parent_b])
    _pay(a, parent_a, "150.00", Payment.Status.SUCCESS, "REF-A-150")
    _pay(a, parent_a, "500.00", Payment.Status.PENDING, "REF-A-500")
    _pay(a, parent_a, "300.00", Payment.Status.FAILED, "REF-A-300")
    _pay(b, parent_b, "999.00", Payment.Status.SUCCESS, "REF-B-999")
    return parent_a, parent_b, a, b


def _refs(rows):
    return [next(t for t in r.split(",") if t.startswith("REF-")) for r in rows]


class TestPaymentsOrdering:
    @pytest.mark.parametrize(
        "ordering, expected",
        [
            ("", ["REF-A-300", "REF-A-500", "REF-A-150"]),
            ("-amount", ["REF-A-500", "REF-A-300", "REF-A-150"]),
            ("amount", ["REF-A-150", "REF-A-300", "REF-A-500"]),
            ("status", ["REF-A-300", "REF-A-500", "REF-A-150"]),
            ("date", ["REF-A-150", "REF-A-500", "REF-A-300"]),
        ],
    )
    def test_list_and_export_share_the_order(self, api_client, two_families, ordering, expected):
        parent_a, *_ = two_families
        api_client.force_authenticate(parent_a)
        params = {"ordering": ordering} if ordering else {}
        rows = api_client.get(reverse("payment-history"), params).data["results"]
        assert [r["gateway_tx_id"] for r in rows] == expected
        exported = _csv_rows(api_client.get(reverse("payment-history-export"), params))
        assert _refs(exported) == expected

    def test_portal_sort_keys_are_whitelisted(self):
        # PaymentsPage.tsx sorts by date / amount / status (gateway stays admin-only UI).
        assert {"date", "amount", "status"} <= set(payment_views.PAYMENT_ORDERING)


class TestPaymentsExport:
    def test_family_scoped_with_two_families(self, api_client, two_families):
        parent_a, parent_b, a, b = two_families
        api_client.force_authenticate(parent_a)
        refs = _refs(_csv_rows(api_client.get(reverse("payment-history-export"))))
        assert sorted(refs) == ["REF-A-150", "REF-A-300", "REF-A-500"]
        # Asking for the other family's child yields nothing, never their rows.
        assert _csv_rows(api_client.get(reverse("payment-history-export"), {"student": b.id})) == []

        api_client.force_authenticate(parent_b)
        assert _refs(_csv_rows(api_client.get(reverse("payment-history-export")))) == ["REF-B-999"]

    def test_filters_are_honoured(self, api_client, two_families):
        parent_a, *_ = two_families
        api_client.force_authenticate(parent_a)
        rows = _csv_rows(api_client.get(reverse("payment-history-export"), {"status": "pending"}))
        assert _refs(rows) == ["REF-A-500"]
        assert (
            _csv_rows(api_client.get(reverse("payment-history-export"), {"from": "2999-01-01"}))
            == []
        )

    def test_header_is_unchanged(self, api_client, two_families):
        parent_a, *_ = two_families
        api_client.force_authenticate(parent_a)
        resp = api_client.get(reverse("payment-history-export"))
        header = b"".join(resp.streaming_content).decode("utf-8-sig").split("\r\n")[0]
        assert header == "Fecha,Alumno,Concepto,Monto,Moneda,Estado,Pasarela,Referencia"

    def test_xlsx(self, api_client, two_families):
        from openpyxl import load_workbook

        parent_a, *_ = two_families
        api_client.force_authenticate(parent_a)
        resp = api_client.get(
            reverse("payment-history-export"), {"fmt": "xlsx", "ordering": "-amount"}
        )
        assert resp.status_code == 200
        assert resp["Content-Disposition"].startswith('attachment; filename="pagos_')
        rows = list(load_workbook(io.BytesIO(resp.content)).active.iter_rows(values_only=True))
        assert rows[0][0] == "Fecha" and len(rows) == 4
        assert float(rows[1][3]) == 500.0

    def test_audited(self, api_client, two_families):
        parent_a, *_ = two_families
        api_client.force_authenticate(parent_a)
        api_client.get(reverse("payment-history-export"), {"status": "success", "ordering": "date"})
        entry = AuditLog.objects.get(object_type="export:payments.family")
        assert entry.actor == parent_a and entry.context == "portal"
        assert entry.changes["rows"] == 1
        assert entry.changes["filters"] == {"status": "success", "ordering": "date"}

    def test_over_cap_is_413_instead_of_silent_truncation(
        self, api_client, two_families, monkeypatch
    ):
        monkeypatch.setattr(FAMILY_PAYMENTS_EXPORT_SPEC, "row_cap", 2)
        parent_a, *_ = two_families
        api_client.force_authenticate(parent_a)
        resp = api_client.get(reverse("payment-history-export"))
        assert resp.status_code == 413
        assert resp.data["detail"] == "Acote los filtros: el límite es 2 filas."

    def test_pdf_not_offered_and_anonymous_refused(self, api_client, two_families):
        assert api_client.get(reverse("payment-history-export")).status_code == 401
        parent_a, *_ = two_families
        api_client.force_authenticate(parent_a)
        assert api_client.get(reverse("payment-history-export"), {"fmt": "pdf"}).status_code == 400

    def test_throttle_scope(self):
        assert payment_views.PaymentHistoryExportView.throttle_scope == "portal-export"


# family student-id resolve + count + rows (topup/student/user joined) + audit insert.
EXPORT_BUDGET = 4


@pytest.mark.parametrize("n_payments", [2, 6])
def test_export_query_budget_is_constant(api_client, django_assert_num_queries, n_payments):
    parent = ParentFactory()
    student = StudentProfileFactory(parents=[parent])
    for i in range(n_payments):
        _pay(student, parent, "100.00", Payment.Status.SUCCESS, f"REF-Q-{i}")
    api_client.force_authenticate(parent)
    with django_assert_num_queries(EXPORT_BUDGET):
        resp = api_client.get(reverse("payment-history-export"))
        assert len(_csv_rows(resp)) == n_payments
