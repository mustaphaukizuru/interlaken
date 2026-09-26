"""
core/importing.py — CSV / XLSX import with dry-run, dedupe and error report (Data Ops C4).

One ``ImportSpec`` per entity describes the columns (header aliases, parsers,
validators), the in-file dedupe keys, how a row matches an existing database
row and how to create/update it. The pipeline is the same for every entity:

    parse_upload(file)  →  validate(rows, spec)  →  commit(report, spec, actor)

* ``parse_upload``: CSV (``utf-8-sig`` then ``latin-1``; delimiter sniffed among
  ``, ; \\t``) or XLSX (openpyxl, read-only, first sheet). Headers are
  normalised (lowercase, accents stripped, spaces → ``_``) and matched against
  each column's aliases, so "Matrícula", "matricula" and "MATRICULA " are the
  same column.
* ``validate``: per-row ``{line, key, action, errors, warnings, data}`` where
  ``action`` is ``crear`` / ``actualizar`` / ``omitir`` / ``error``. In-file
  duplicates mark the LATER row ``Duplicado de la fila N``; a DB match decides
  crear vs actualizar. Everything the commit would check runs here, so the
  dry-run preview is what the commit will do.
* ``commit``: outer ``transaction.atomic()``, one savepoint per row, a broad
  ``except`` isolating each row, ``record()`` per created/updated instance
  (``context='import:<entity>'``) and one ``record_import`` summary.
* ``render_report``: the original columns + ``fila, resultado, errores, avisos``
  as CSV or XLSX (rows with errors first). Stateless: the client re-posts the
  same file with ``report=csv|xlsx`` to download it.

Caps (no workers, one request): 2,000 rows / 5 MB, answered with HTTP 413.
"""

from __future__ import annotations

import csv
import io
import logging
import re
import unicodedata
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from django.db import transaction
from django.http import HttpResponse
from django.utils.dateparse import parse_date
from rest_framework import status
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from .audit import record, record_import
from .exceptions import PayloadTooLarge, cap_message, error_response
from .exporting import CONTENT_TYPES, xlsx_bytes
from .exports import as_download, export_filename
from .permissions import IsAdmin
from .throttling import SharedScopedRateThrottle

logger = logging.getLogger(__name__)

MAX_ROWS = 2_000
MAX_BYTES = 5 * 1024 * 1024
ACTIONS = ("crear", "actualizar", "omitir", "error")
REPORT_COLUMNS = ("fila", "resultado", "errores", "avisos")
TRUE_WORDS = frozenset({"1", "true", "si", "sí", "x", "verdadero", "yes", "y", "s"})
FALSE_WORDS = frozenset({"", "0", "false", "no", "falso", "n"})


# ── header normalisation ──────────────────────────────────
def normalize_header(text) -> str:
    """``'  Matrícula del Alumno '`` → ``'matricula_del_alumno'``."""
    folded = (
        "".join(
            ch
            for ch in unicodedata.normalize("NFKD", str(text or ""))
            if not unicodedata.combining(ch)
        )
        .lower()
        .strip()
    )
    folded = re.sub(r"[\s\-/.]+", "_", folded)
    folded = re.sub(r"[^a-z0-9_]", "", folded)
    return re.sub(r"_+", "_", folded).strip("_")


# ── parsers (reusable in specs) ───────────────────────────
def parse_text(value: str) -> str:
    return str(value or "").strip()


def parse_int(value: str) -> int:
    try:
        return int(Decimal(str(value).strip().replace(",", "")))
    except (InvalidOperation, ValueError):
        raise ValueError("debe ser un número entero") from None


def parse_decimal(value: str) -> Decimal:
    text = str(value).strip().replace("$", "").replace(" ", "").replace(",", "")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        raise ValueError("debe ser un monto (por ejemplo 150.00)") from None


def parse_date_es(value: str) -> date:
    """``DD/MM/YYYY``, ``DD-MM-YYYY`` or ISO ``YYYY-MM-DD``."""
    text = str(value).strip()
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$", text)
    if match:
        d, m, y = (int(p) for p in match.groups())
        try:
            return date(y, m, d)
        except ValueError:
            raise ValueError("fecha inválida; use DD/MM/AAAA") from None
    parsed = parse_date(text[:10])
    if parsed is None:
        raise ValueError("fecha inválida; use DD/MM/AAAA")
    return parsed


def parse_bool(value: str) -> bool:
    text = str(value or "").strip().lower()
    if text in TRUE_WORDS:
        return True
    if text in FALSE_WORDS:
        return False
    raise ValueError("use sí o no")


def parse_email(value: str) -> str:
    text = str(value or "").strip().lower()
    if text and ("@" not in text or " " in text):
        raise ValueError("correo inválido")
    return text


def parse_choice(choices, *, aliases: dict[str, str] | None = None):
    """Parser factory: value must be one of ``choices`` (labels and aliases accepted)."""
    lookup = {}
    for item in choices:
        value, label = (item[0], item[1]) if isinstance(item, (tuple, list)) else (item, item)
        lookup[normalize_header(value)] = value
        lookup[normalize_header(label)] = value
    for alias, value in (aliases or {}).items():
        lookup[normalize_header(alias)] = value
    allowed = ", ".join(str(v) for v in dict.fromkeys(lookup.values()))

    def parse(value: str):
        key = normalize_header(value)
        if key in lookup:
            return lookup[key]
        raise ValueError(f"use uno de: {allowed}")

    return parse


# ── spec ──────────────────────────────────────────────────
@dataclass(frozen=True)
class ImportCol:
    key: str
    header_aliases: tuple[str, ...] = ()
    required: bool = False
    parse: Callable[[str], Any] | None = None
    validators: tuple[Callable[[Any, dict], str | None], ...] = ()
    example: str = ""
    label: str = ""

    @property
    def header(self) -> str:
        return self.label or self.key

    def matches(self, normalized: str) -> bool:
        candidates = {normalize_header(self.key), normalize_header(self.header)}
        candidates.update(normalize_header(a) for a in self.header_aliases)
        return normalized in candidates


@dataclass
class ImportSpec:
    entity: str
    columns: list[ImportCol]
    dedupe_keys: list[tuple[str, ...]] = field(default_factory=list)
    db_match: Callable[[dict], Any] | None = None
    create: Callable[[dict, Any], Any] | None = None
    update: Callable[[Any, dict, Any], Any] | None = None
    validate_row: Callable[[dict, ImportRow], None] | None = None
    skip_reason: Callable[[dict, Any], str | None] | None = None
    key_of: Callable[[dict], str] | None = None
    max_rows: int = MAX_ROWS
    max_bytes: int = MAX_BYTES
    duplicates_block: bool = True
    label: str = ""

    def headers(self) -> list[str]:
        return [c.header for c in self.columns]

    def example_row(self) -> list[str]:
        return [c.example for c in self.columns]

    def column(self, key: str) -> ImportCol | None:
        return next((c for c in self.columns if c.key == key), None)


@dataclass
class ImportRow:
    line: int
    key: str = ""
    action: str = "crear"
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    data: dict = field(default_factory=dict)
    raw: dict = field(default_factory=dict)
    instance: Any = None

    def error(self, message: str) -> None:
        self.errors.append(message)
        self.action = "error"

    def warn(self, message: str) -> None:
        self.warnings.append(message)

    def as_dict(self) -> dict:
        return {
            "line": self.line,
            "key": self.key,
            "action": self.action,
            "errors": list(self.errors),
            "warnings": list(self.warnings),
            "data": {k: _jsonable(v) for k, v in self.data.items()},
        }


@dataclass
class ImportReport:
    entity: str
    rows: list[ImportRow]
    headers: list[str]
    fmt: str = "csv"
    dry_run: bool = True

    @property
    def counts(self) -> dict[str, int]:
        out = dict.fromkeys(ACTIONS, 0)
        for row in self.rows:
            out[row.action] = out.get(row.action, 0) + 1
        return out

    @property
    def has_errors(self) -> bool:
        return any(r.action == "error" for r in self.rows)

    def as_dict(self) -> dict:
        return {
            "entity": self.entity,
            "dry_run": self.dry_run,
            "total_rows": len(self.rows),
            "counts": self.counts,
            "rows": [r.as_dict() for r in self.rows],
        }


class UploadError(Exception):
    """A 400 for a file the pipeline cannot use (empty, unreadable, headers missing)."""

    def __init__(self, message: str, **extra):
        super().__init__(message)
        self.message = message
        self.extra = extra


def _jsonable(value):
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


# ── parsing ───────────────────────────────────────────────
def _sniff_delimiter(line: str) -> str:
    counts = {d: line.count(d) for d in (",", ";", "\t")}
    best = max(counts, key=counts.get)
    return best if counts[best] else ","


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "sí" if value else "no"
    if isinstance(value, datetime):
        return (
            value.date().isoformat()
            if not (value.hour or value.minute or value.second)
            else value.isoformat(sep=" ")
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def _parse_csv(data: bytes):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    text = text.lstrip("﻿")
    if not text.strip():
        raise UploadError("Archivo vacío o sin encabezados.")
    first_line = text.splitlines()[0]
    reader = csv.reader(io.StringIO(text), delimiter=_sniff_delimiter(first_line))
    headers = [str(h or "").strip() for h in next(reader, [])]
    while headers and not headers[-1]:
        headers.pop()
    if not any(headers):
        raise UploadError("Archivo vacío o sin encabezados.")
    rows = []
    for values in reader:
        if not any(str(v or "").strip() for v in values):
            continue
        row = {
            h: (str(values[i]).strip() if i < len(values) and values[i] is not None else "")
            for i, h in enumerate(headers)
            if h
        }
        rows.append((reader.line_num, row))
    return rows, headers


def _parse_xlsx(data: bytes):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise UploadError("No se pudo leer el archivo XLSX.") from None
    try:
        ws = wb.worksheets[0]
        iterator = ws.iter_rows(values_only=True)
        headers = [_cell_text(h) for h in next(iterator, ())]
        while headers and not headers[-1]:
            headers.pop()
        if not any(headers):
            raise UploadError("Archivo vacío o sin encabezados.")
        rows = []
        for line, values in enumerate(iterator, start=2):
            cells = [_cell_text(v) for v in values]
            if not any(cells):
                continue
            row = {h: (cells[i] if i < len(cells) else "") for i, h in enumerate(headers) if h}
            rows.append((line, row))
    finally:
        wb.close()
    return rows, headers


def parse_upload(upload, *, max_bytes: int = MAX_BYTES):
    """→ ``(rows, headers, fmt)`` where ``rows`` is ``[(line, {header: text}), ...]``.

    Raises ``PayloadTooLarge`` above ``max_bytes`` and ``UploadError`` for an
    unusable file. The format is taken from the name (``.xlsx``) or the zip
    magic bytes, so a mislabelled Excel file still parses.
    """
    name = (getattr(upload, "name", "") or "").lower()
    data = upload.read() if hasattr(upload, "read") else bytes(upload)
    if len(data) > max_bytes:
        raise PayloadTooLarge(
            detail=f"El archivo excede el límite de {max_bytes // (1024 * 1024)} MB."
        )
    if name.endswith(".xlsx") or data[:4] == b"PK\x03\x04":
        rows, headers = _parse_xlsx(data)
        return rows, headers, "xlsx"
    rows, headers = _parse_csv(data)
    return rows, headers, "csv"


def match_headers(headers, spec: ImportSpec) -> dict[str, str]:
    """``{original header: column key}``; 400 when a required column is missing."""
    mapping: dict[str, str] = {}
    for header in headers:
        normalized = normalize_header(header)
        if not normalized:
            continue
        for col in spec.columns:
            if col.key in mapping.values():
                continue
            if col.matches(normalized):
                mapping[header] = col.key
                break
    missing = [c.header for c in spec.columns if c.required and c.key not in mapping.values()]
    if missing:
        raise UploadError(
            f'Faltan columnas requeridas: {", ".join(missing)}.',
            missing=missing,
            expected_headers=spec.headers(),
        )
    return mapping


# ── validation ────────────────────────────────────────────
def _dedupe_value(value) -> str:
    return str(value if value is not None else "").strip().lower()


def validate(rows, spec: ImportSpec, *, headers=None, fmt: str = "csv") -> ImportReport:
    """Parse, validate, dedupe (in file + DB) → ``ImportReport`` (no writes)."""
    if headers is None:
        headers = list(dict.fromkeys(h for _, row in rows for h in row))
    mapping = match_headers(headers, spec)
    header_for = {key: header for header, key in mapping.items()}

    report_rows: list[ImportRow] = []
    for line, raw in rows:
        row = ImportRow(line=line, raw=dict(raw))
        data: dict[str, Any] = {}
        for col in spec.columns:
            text = raw.get(header_for.get(col.key, ""), "")
            text = str(text or "").strip()
            if not text:
                if col.required:
                    row.error(f"Falta {col.header}.")
                    continue
                # Empty optional cell: '' for free text, None for a typed column.
                data[col.key] = "" if col.parse is None else None
                continue
            try:
                data[col.key] = col.parse(text) if col.parse else text
            except (ValueError, InvalidOperation, TypeError) as exc:
                row.error(f'{col.header}: {exc or "valor no válido"}.')
                continue
            for check in col.validators:
                message = check(data[col.key], data)
                if message:
                    row.error(f"{col.header}: {message}.")
        row.data = data
        row.key = _row_key(spec, data)
        if spec.validate_row is not None:
            spec.validate_row(data, row)
        report_rows.append(row)

    # In-file duplicates: the LATER row is flagged, the first one stands.
    for key_fields in spec.dedupe_keys:
        seen: dict[tuple, int] = {}
        for row in report_rows:
            values = tuple(_dedupe_value(row.data.get(f)) for f in key_fields)
            if not all(values):
                continue
            first = seen.get(values)
            if first is None:
                seen[values] = row.line
                continue
            message = f"Duplicado de la fila {first}."
            if spec.duplicates_block:
                row.error(message)
            else:
                row.warn(message)

    # DB match decides crear vs actualizar (even for error rows: the key is useful).
    for row in report_rows:
        if spec.db_match is None:
            continue
        try:
            row.instance = spec.db_match(row.data)
        except Exception as exc:  # a broken key must not 500 the preview
            row.error(f"No se pudo verificar en la base de datos: {exc}.")
            continue
        if row.action == "error":
            continue
        row.action = "actualizar" if row.instance is not None else "crear"
        if spec.skip_reason is not None:
            reason = spec.skip_reason(row.data, row.instance)
            if reason:
                row.action = "omitir"
                row.warn(reason)

    return ImportReport(entity=spec.entity, rows=report_rows, headers=list(headers), fmt=fmt)


def _row_key(spec: ImportSpec, data: dict) -> str:
    if spec.key_of is not None:
        try:
            return str(spec.key_of(data) or "")
        except Exception:
            return ""
    for key_fields in spec.dedupe_keys:
        values = [str(data.get(f) or "") for f in key_fields]
        if all(values):
            return " / ".join(values)
    for col in spec.columns:
        if col.required and data.get(col.key):
            return str(data[col.key])
    return ""


# ── commit ────────────────────────────────────────────────
def commit(report: ImportReport, spec: ImportSpec, actor, *, filename: str = "") -> ImportReport:
    """Apply every ``crear``/``actualizar`` row; each row in its own savepoint."""
    context = f"import:{spec.entity}"
    with transaction.atomic():
        for row in report.rows:
            if row.action not in ("crear", "actualizar"):
                continue
            try:
                with transaction.atomic():
                    if row.action == "crear":
                        if spec.create is None:
                            raise RuntimeError("esta importación no crea filas nuevas")
                        instance = spec.create(row.data, actor)
                        audit_action = "create"
                    else:
                        if spec.update is None:
                            raise RuntimeError("esta importación no actualiza filas")
                        instance = spec.update(row.instance, row.data, actor)
                        audit_action = "update"
                    if instance is not None:
                        record(
                            audit_action,
                            instance,
                            {"import": spec.entity, "line": row.line, "key": row.key},
                            actor=actor,
                            context=context,
                        )
                    row.instance = instance
            except Exception as exc:  # per-row isolation: report, keep importing
                logger.warning("import %s line %s failed: %s", spec.entity, row.line, exc)
                row.error(f"Error inesperado: {_message(exc)}")
        record_import(spec.entity, report.counts, actor, dry_run=False, filename=filename)
    report.dry_run = False
    return report


def _message(exc: Exception) -> str:
    detail = getattr(exc, "detail", None)
    if detail is not None:
        if isinstance(detail, dict):
            return "; ".join(f"{k}: {_flatten(v)}" for k, v in detail.items())
        return _flatten(detail)
    messages = getattr(exc, "messages", None)
    if messages:
        return "; ".join(str(m) for m in messages)
    return str(exc) or exc.__class__.__name__


def _flatten(value) -> str:
    if isinstance(value, (list, tuple)):
        return "; ".join(_flatten(v) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}: {_flatten(v)}" for k, v in value.items())
    return str(value)


# ── files: report + template ──────────────────────────────
def report_rows(report: ImportReport) -> tuple[list[str], list[list[str]]]:
    headers = list(report.headers) + list(REPORT_COLUMNS)
    ordered = sorted(report.rows, key=lambda r: 0 if r.action == "error" else 1)
    rows = [
        [r.raw.get(h, "") for h in report.headers]
        + [str(r.line), r.action, " | ".join(r.errors), " | ".join(r.warnings)]
        for r in ordered
    ]
    return headers, rows


def _csv_response(headers, rows, filename: str) -> HttpResponse:
    buf = io.StringIO()
    buf.write("﻿")
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    response = HttpResponse(buf.getvalue().encode("utf-8"), content_type=CONTENT_TYPES["csv"])
    return as_download(response, filename)


def _xlsx_response(headers, rows, filename: str, *, sheet_title: str) -> HttpResponse:
    response = HttpResponse(
        xlsx_bytes(headers, rows, sheet_title=sheet_title), content_type=CONTENT_TYPES["xlsx"]
    )
    return as_download(response, filename)


def render_report(report: ImportReport, fmt: str | None = None) -> HttpResponse:
    """The uploaded rows + ``fila, resultado, errores, avisos``, errors first."""
    fmt = (fmt or report.fmt or "csv").lower()
    headers, rows = report_rows(report)
    filename = export_filename(f"reporte_{report.entity}", fmt)
    if fmt == "xlsx":
        return _xlsx_response(headers, rows, filename, sheet_title="Reporte")
    return _csv_response(headers, rows, filename)


def render_template(spec: ImportSpec, fmt: str = "csv") -> HttpResponse:
    """Headers + one example row, as CSV or XLSX."""
    fmt = (fmt or "csv").lower()
    filename = export_filename(f"plantilla_{spec.entity}", fmt)
    if fmt == "xlsx":
        return _xlsx_response(
            spec.headers(), [spec.example_row()], filename, sheet_title="Plantilla"
        )
    return _csv_response(spec.headers(), [spec.example_row()], filename)


# ── views ─────────────────────────────────────────────────
def truthy(value, default: bool) -> bool:
    if value is None:
        return default
    return str(value).strip().lower() in TRUE_WORDS


class ImportTemplateView(APIView):
    """``GET .../import/template/?fmt=csv|xlsx`` — headers + one example row."""

    permission_classes = [IsAdmin]
    spec: ImportSpec | None = None

    def get_spec(self) -> ImportSpec:
        if self.spec is None:
            raise NotImplementedError("spec is required")
        return self.spec

    def get(self, request):
        fmt = (request.query_params.get("fmt") or "csv").lower()
        if fmt not in ("csv", "xlsx"):
            return error_response("Use fmt=csv o fmt=xlsx.")
        return render_template(self.get_spec(), fmt)


class ImportView(APIView):
    """``POST .../import/`` multipart ``file`` + ``dry_run=1|0`` (+ ``report=csv|xlsx``).

    Stateless: the dry-run and the commit both re-parse the file. ``report``
    returns the annotated file of the dry-run instead of JSON. A commit with
    error rows is refused unless ``only_valid=1`` (the UI sends ``valid_only=1``;
    both spellings are accepted) — "Importar solo las filas válidas" imports the
    crear/actualizar rows and leaves the rest.
    """

    permission_classes = [IsAdmin]
    throttle_classes = [SharedScopedRateThrottle]
    throttle_scope = "admin-import"
    parser_classes = [MultiPartParser, FormParser]
    spec: ImportSpec | None = None
    template_url: str = ""

    def get_spec(self) -> ImportSpec:
        if self.spec is None:
            raise NotImplementedError("spec is required")
        return self.spec

    def get_template_url(self) -> str:
        return self.template_url

    def get_actor(self):
        return self.request.user

    def post(self, request):
        spec = self.get_spec()
        upload = request.FILES.get("file")
        if upload is None:
            return error_response(
                'Adjunte un archivo CSV o XLSX en el campo "file".',
                expected_headers=spec.headers(),
                template=self.get_template_url(),
            )
        if getattr(upload, "size", 0) > spec.max_bytes:
            raise PayloadTooLarge(
                detail=f"El archivo excede el límite de {spec.max_bytes // (1024 * 1024)} MB."
            )
        try:
            rows, headers, fmt = parse_upload(upload, max_bytes=spec.max_bytes)
            if len(rows) > spec.max_rows:
                raise PayloadTooLarge(detail=cap_message(spec.max_rows, verb="Divida el archivo"))
            report = validate(rows, spec, headers=headers, fmt=fmt)
        except UploadError as exc:
            return error_response(
                exc.message,
                status.HTTP_400_BAD_REQUEST,
                template=self.get_template_url(),
                **exc.extra,
            )

        report_fmt = (request.data.get("report") or "").lower()
        if report_fmt in ("csv", "xlsx"):
            return render_report(report, report_fmt)

        if truthy(request.data.get("dry_run"), True):
            return Response(report.as_dict())

        only_valid = truthy(request.data.get("only_valid"), False) or truthy(
            request.data.get("valid_only"), False
        )
        if report.has_errors and not only_valid:
            counts = report.counts
            return error_response(
                f'El archivo tiene {counts["error"]} fila(s) con error. Corríjalas o marque '
                f"«Importar solo las filas válidas».",
                counts=counts,
            )

        commit(report, spec, self.get_actor(), filename=getattr(upload, "name", "") or "")
        return Response(report.as_dict())
