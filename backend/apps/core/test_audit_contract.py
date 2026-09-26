"""
Data Ops Phase 7: the Auditoría list on the list contract and its export.

``ordering=actor`` sorts by the actor's email; ``q`` searches actor label,
context and object id; ``object_type`` / ``object_id`` / ``context`` filters;
the export is the list view (same filters and sort) in CSV / XLSX / PDF and is
itself audited.
"""

import io

import pytest
from django.urls import reverse

from apps.accounts.factories import AdminFactory
from apps.core.audit import record
from apps.core.filters import AUDIT_ORDERING
from apps.core.models import AuditLog

pytestmark = pytest.mark.django_db

LIST = reverse("core-admin-audit")
EXPORT = reverse("core-admin-audit-export")

# Mirrors AUDIT_ORDERING_KEYS in frontend/src/hooks/queries/audit.ts.
FRONTEND_ORDERING_KEYS = {"date", "actor", "action", "object", "context"}


@pytest.fixture
def entries():
    zed = AdminFactory(email="zed@test.mx")
    ada = AdminFactory(email="ada@test.mx")
    # Every assertion scopes to the demo.* rows: factories emit their own
    # signal-audited rows (role changes) that must not interfere.
    rows = [
        record(
            "update",
            object_type="demo.thing",
            object_id="7",
            changes={"status": ["a", "b"]},
            actor=zed,
            context="demo.alpha",
        ),
        record(
            "update",
            object_type="demo.thing",
            object_id="8",
            changes={"reason": "x"},
            actor=ada,
            context="demo.beta",
        ),
        record(
            "create",
            object_type="demo.other",
            object_id="7",
            changes={},
            actor_label="system:cron",
            context="demo.gamma",
        ),
    ]
    return {"rows": rows, "zed": zed, "ada": ada}


def _ids(resp):
    return [r["id"] for r in resp.data["results"]]


def test_ordering_keys_match_the_frontend_constant():
    assert set(AUDIT_ORDERING) == FRONTEND_ORDERING_KEYS


class TestList:
    def test_ordering_actor_sorts_by_actor_email(self, admin_client, entries):
        r = entries["rows"]
        ids = _ids(
            admin_client.get(LIST, {"ordering": "actor", "context": "demo", "page_size": 100})
        )
        # NULL actor (system) sorts first on SQLite ascending; ada before zed.
        assert ids.index(r[1].id) < ids.index(r[0].id)
        desc = _ids(
            admin_client.get(LIST, {"ordering": "-actor", "context": "demo", "page_size": 100})
        )
        assert desc.index(r[0].id) < desc.index(r[1].id)

    def test_q_over_actor_label_context_and_object_id(self, admin_client, entries):
        r = entries["rows"]
        assert _ids(admin_client.get(LIST, {"q": "demo.beta"})) == [r[1].id]
        assert _ids(admin_client.get(LIST, {"q": "system:cron"})) == [r[2].id]
        assert set(_ids(admin_client.get(LIST, {"q": "7", "context": "demo"}))) == {
            r[0].id,
            r[2].id,
        }

    def test_object_filters(self, admin_client, entries):
        r = entries["rows"]
        resp = admin_client.get(LIST, {"object_type": "demo.thing", "object_id": "7"})
        assert _ids(resp) == [r[0].id]
        assert _ids(admin_client.get(LIST, {"object_type": "demo.other"})) == [r[2].id]

    def test_unknown_action_is_ignored(self, admin_client, entries):
        assert admin_client.get(LIST, {"action": "nope"}).status_code == 200

    def test_export_action_is_filterable(self, admin_client, entries):
        admin_client.get(EXPORT, {"fmt": "csv", "context": "demo"})
        resp = admin_client.get(LIST, {"action": "export"})
        assert resp.data["count"] == 1 and resp.data["results"][0]["object_type"] == "export:audit"

    @pytest.mark.parametrize("n", [3, 12])
    def test_query_budget_is_constant(self, api_client, django_assert_num_queries, n):
        admin = AdminFactory()
        for i in range(n):
            record("update", object_type="demo.bulk", object_id=str(i), actor=admin, context="demo")
        api_client.force_authenticate(admin)
        with django_assert_num_queries(2):
            resp = api_client.get(LIST, {"q": "demo", "ordering": "-actor", "page_size": 50})
        assert resp.status_code == 200


class TestExport:
    def test_csv_honours_filters_and_sort(self, admin_client, entries):
        resp = admin_client.get(
            EXPORT, {"fmt": "csv", "object_type": "demo.thing", "ordering": "actor"}
        )
        assert resp.status_code == 200
        lines = b"".join(resp.streaming_content).decode("utf-8-sig").strip().split("\r\n")
        assert lines[0] == "Fecha,Actor,Acción,Objeto,ID,Contexto,Cambios"
        assert len(lines) == 3
        assert "ada@test.mx" in lines[1] and "zed@test.mx" in lines[2]

    def test_xlsx_and_pdf(self, admin_client, entries):
        from openpyxl import load_workbook

        resp = admin_client.get(EXPORT, {"fmt": "xlsx", "context": "demo"})
        ws = load_workbook(io.BytesIO(resp.content)).active
        assert ws.max_row == 4
        pdf = admin_client.get(EXPORT, {"fmt": "pdf", "context": "demo"})
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")

    def test_export_is_audited(self, admin_client, entries):
        admin_client.get(EXPORT, {"fmt": "xlsx", "context": "demo"})
        entry = AuditLog.objects.get(object_type="export:audit")
        assert entry.object_id == "xlsx" and entry.changes["rows"] == 3
