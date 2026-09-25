"""
apps.core.importing — parsers (CSV utf-8-sig / latin-1 / ; / tab, XLSX), header
aliases, in-file + DB dedupe, per-row savepoint commit with audit, report and
template rendering, and the 2,000 rows / 5 MB caps → 413.
"""

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.accounts.factories import AdminFactory
from apps.core.exceptions import PayloadTooLarge, cap_message
from apps.core.importing import (
    ImportCol,
    ImportSpec,
    ImportTemplateView,
    ImportView,
    UploadError,
    commit,
    match_headers,
    normalize_header,
    parse_bool,
    parse_choice,
    parse_date_es,
    parse_decimal,
    parse_email,
    parse_int,
    parse_upload,
    render_report,
    render_template,
    validate,
)
from apps.core.models import AuditLog, ContactMessage

pytestmark = pytest.mark.django_db


def _create(data, actor):
    if data["nombre"] == "BOOM":
        raise RuntimeError("fila envenenada")
    return ContactMessage.objects.create(
        name=data["nombre"],
        email=data["correo"],
        subject=data.get("asunto") or "-",
        message=str(data.get("monto") or ""),
    )


def _update(instance, data, actor):
    instance.name = data["nombre"]
    instance.save(update_fields=["name"])
    return instance


SPEC = ImportSpec(
    entity="mensajes",
    columns=[
        ImportCol(
            "nombre",
            ("name", "nombre completo"),
            required=True,
            label="Nombre",
            example="Ana Pérez",
        ),
        ImportCol(
            "correo",
            ("email", "e-mail"),
            required=True,
            parse=parse_email,
            label="Correo",
            example="ana@x.mx",
        ),
        ImportCol("asunto", label="Asunto", example="Informes"),
        ImportCol("monto", parse=parse_decimal, label="Monto", example="150.00"),
        ImportCol("fecha", parse=parse_date_es, label="Fecha", example="05/03/2026"),
        ImportCol("activo", parse=parse_bool, label="Activo", example="sí"),
    ],
    dedupe_keys=[("correo",)],
    db_match=lambda d: (
        ContactMessage.objects.filter(email=d["correo"]).first() if d.get("correo") else None
    ),
    create=_create,
    update=_update,
)


def _upload(content, name="datos.csv"):
    f = io.BytesIO(content if isinstance(content, bytes) else content.encode("utf-8"))
    f.name = name
    f.size = len(f.getvalue())
    return f


def _xlsx_bytes(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestParsers:
    def test_normalize_header(self):
        assert normalize_header("  Matrícula del Alumno ") == "matricula_del_alumno"
        assert normalize_header("E-mail") == "e_mail"
        assert normalize_header("Nombre/Apellidos") == "nombre_apellidos"
        assert normalize_header(None) == ""

    def test_value_parsers(self):
        assert parse_int("1,200") == 1200
        assert parse_decimal("$1,234.5") == Decimal("1234.50")
        assert parse_date_es("05/03/2026") == date(2026, 3, 5)
        assert parse_date_es("2026-03-05") == date(2026, 3, 5)
        assert parse_bool("Sí") is True and parse_bool("no") is False and parse_bool("") is False
        assert parse_email(" Ana@X.mx ") == "ana@x.mx"
        choice = parse_choice(
            [("preescolar", "Preescolar"), ("primaria", "Primaria")],
            aliases={"kinder": "preescolar"},
        )
        assert choice("Kínder") == "preescolar" and choice("PRIMARIA") == "primaria"
        for fn, bad in (
            (parse_int, "x"),
            (parse_decimal, "abc"),
            (parse_date_es, "31/31/2026"),
            (parse_bool, "tal vez"),
            (parse_email, "sin arroba"),
            (choice, "secundaria"),
        ):
            with pytest.raises(ValueError):
                fn(bad)

    def test_csv_utf8_sig_with_accents(self):
        rows, headers, fmt = parse_upload(_upload("﻿Nombre,Correo\nJosé Núñez,jose@x.mx\n"))
        assert fmt == "csv" and headers == ["Nombre", "Correo"]
        assert rows == [(2, {"Nombre": "José Núñez", "Correo": "jose@x.mx"})]

    def test_csv_latin1_fallback(self):
        raw = "Nombre;Correo\nJosé;jose@x.mx\n".encode("latin-1")
        rows, headers, _ = parse_upload(_upload(raw))
        assert headers == ["Nombre", "Correo"]
        assert rows[0][1]["Nombre"] == "José"

    def test_csv_delimiters_are_sniffed(self):
        for sep in (";", "\t", ","):
            rows, headers, _ = parse_upload(_upload(f"a{sep}b\n1{sep}2\n\n\n"))
            assert headers == ["a", "b"] and rows == [(2, {"a": "1", "b": "2"})]

    def test_xlsx_first_sheet_with_typed_cells(self):
        raw = _xlsx_bytes(
            [
                ["Nombre", "Monto", "Entero", "Fecha", None],
                ["Ana", 12.5, 5.0, date(2026, 3, 5), None],
                [None, None, None, None, None],
                ["Luis", 7, 3, None, None],
            ]
        )
        rows, headers, fmt = parse_upload(_upload(raw, name="datos.xlsx"))
        assert fmt == "xlsx" and headers == ["Nombre", "Monto", "Entero", "Fecha"]
        assert rows[0] == (
            2,
            {"Nombre": "Ana", "Monto": "12.5", "Entero": "5", "Fecha": "2026-03-05"},
        )
        assert rows[1][0] == 4 and rows[1][1]["Nombre"] == "Luis"

    def test_xlsx_detected_by_magic_bytes_when_misnamed(self):
        rows, _, fmt = parse_upload(_upload(_xlsx_bytes([["a"], ["1"]]), name="datos.csv"))
        assert fmt == "xlsx" and rows == [(2, {"a": "1"})]

    def test_empty_or_oversized_files(self):
        with pytest.raises(UploadError):
            parse_upload(_upload("   \n"))
        with pytest.raises(PayloadTooLarge):
            parse_upload(_upload("a,b\n1,2\n"), max_bytes=3)


class TestHeadersAndValidate:
    def test_aliases_and_accents_match_headers(self):
        mapping = match_headers(["NOMBRE COMPLETO", "E-mail", "Asunto", "Otra"], SPEC)
        assert mapping == {"NOMBRE COMPLETO": "nombre", "E-mail": "correo", "Asunto": "asunto"}

    def test_missing_required_headers(self):
        with pytest.raises(UploadError) as exc:
            match_headers(["Nombre"], SPEC)
        assert "Correo" in exc.value.message
        assert exc.value.extra["missing"] == ["Correo"]
        assert exc.value.extra["expected_headers"] == SPEC.headers()

    def test_validate_reports_per_row(self):
        existing = ContactMessage.objects.create(
            name="Old", email="ana@x.mx", subject="s", message="m"
        )
        rows = [
            (2, {"Nombre": "Ana", "Correo": "ana@x.mx", "Monto": "10"}),
            (3, {"Nombre": "Beto", "Correo": "beto@x.mx", "Monto": "abc", "Fecha": "99/99/2026"}),
            (4, {"Nombre": "", "Correo": "c@x.mx"}),
            (5, {"Nombre": "Ana bis", "Correo": "ANA@x.mx"}),
            (6, {"Nombre": "Dora", "Correo": "dora@x.mx", "Activo": "sí"}),
        ]
        report = validate(rows, SPEC, headers=["Nombre", "Correo", "Monto", "Fecha", "Activo"])
        by_line = {r.line: r for r in report.rows}
        assert by_line[2].action == "actualizar" and by_line[2].instance == existing
        assert by_line[2].data["monto"] == Decimal("10.00") and by_line[2].key == "ana@x.mx"
        assert by_line[3].action == "error"
        assert by_line[3].errors == [
            "Monto: debe ser un monto (por ejemplo 150.00).",
            "Fecha: fecha inválida; use DD/MM/AAAA.",
        ]
        assert by_line[4].errors == ["Falta Nombre."]
        assert by_line[5].action == "error" and by_line[5].errors == ["Duplicado de la fila 2."]
        assert by_line[6].action == "crear" and by_line[6].data["activo"] is True
        assert by_line[6].data["monto"] is None
        assert report.counts == {"crear": 1, "actualizar": 1, "omitir": 0, "error": 3}
        assert report.has_errors

    def test_duplicates_can_warn_instead_of_block(self):
        spec = ImportSpec("x", SPEC.columns, dedupe_keys=[("correo",)], duplicates_block=False)
        report = validate(
            [(2, {"Nombre": "A", "Correo": "a@x.mx"}), (3, {"Nombre": "B", "Correo": "a@x.mx"})],
            spec,
        )
        assert report.rows[1].action == "crear"
        assert report.rows[1].warnings == ["Duplicado de la fila 2."]

    def test_skip_reason_and_validate_row_hooks(self):
        def skip(data, instance):
            return "Sin cambios." if data["nombre"] == "Skip" else None

        def check(data, row):
            if data.get("asunto") == "raro":
                row.warn("Asunto poco usual.")

        spec = ImportSpec(
            "x", SPEC.columns, db_match=lambda d: None, skip_reason=skip, validate_row=check
        )
        report = validate(
            [
                (2, {"Nombre": "Skip", "Correo": "a@x.mx"}),
                (3, {"Nombre": "Go", "Correo": "b@x.mx", "Asunto": "raro"}),
            ],
            spec,
        )
        assert report.rows[0].action == "omitir" and report.rows[0].warnings == ["Sin cambios."]
        assert report.rows[1].action == "crear" and report.rows[1].warnings == [
            "Asunto poco usual."
        ]


class TestCommit:
    def test_rows_are_isolated_and_audited(self):
        old = ContactMessage.objects.create(name="Old", email="ana@x.mx", subject="s", message="m")
        admin = AdminFactory()
        report = validate(
            [
                (2, {"Nombre": "Ana", "Correo": "ana@x.mx"}),
                (3, {"Nombre": "BOOM", "Correo": "boom@x.mx"}),
                (4, {"Nombre": "Dora", "Correo": "dora@x.mx"}),
                (5, {"Nombre": "", "Correo": "e@x.mx"}),
            ],
            SPEC,
        )
        commit(report, SPEC, admin, filename="datos.csv")
        assert report.dry_run is False
        by_line = {r.line: r for r in report.rows}
        assert by_line[2].action == "actualizar"
        old.refresh_from_db()
        assert old.name == "Ana"
        assert by_line[3].action == "error" and "fila envenenada" in by_line[3].errors[0]
        assert not ContactMessage.objects.filter(email="boom@x.mx").exists()
        assert ContactMessage.objects.filter(email="dora@x.mx").exists()
        assert by_line[5].action == "error"
        created = AuditLog.objects.get(object_type="core.contactmessage", action="create")
        assert created.context == "import:mensajes" and created.changes["line"] == 4
        updated = AuditLog.objects.get(object_type="core.contactmessage", action="update")
        assert updated.object_id == str(old.pk) and updated.actor == admin
        summary = AuditLog.objects.get(object_type="import:mensajes")
        assert summary.action == "import" and summary.object_id == "datos.csv"
        assert summary.changes["counts"] == {"crear": 1, "actualizar": 1, "omitir": 0, "error": 2}


class TestFiles:
    def test_report_has_original_columns_plus_result_errors_first(self):
        report = validate(
            [(2, {"Nombre": "Ana", "Correo": "ana@x.mx"}), (3, {"Nombre": "", "Correo": "b@x.mx"})],
            SPEC,
            headers=["Nombre", "Correo"],
        )
        resp = render_report(report)
        assert resp["Content-Disposition"].startswith('attachment; filename="reporte_mensajes_')
        text = resp.content.decode("utf-8-sig")
        lines = text.split("\r\n")
        assert lines[0] == "Nombre,Correo,fila,resultado,errores,avisos"
        assert lines[1] == ",b@x.mx,3,error,Falta Nombre.,"
        assert lines[2] == "Ana,ana@x.mx,2,crear,,"
        ws = load_workbook(io.BytesIO(render_report(report, "xlsx").content)).worksheets[0]
        assert [c.value for c in ws[1]] == [
            "Nombre",
            "Correo",
            "fila",
            "resultado",
            "errores",
            "avisos",
        ]
        assert ws["D2"].value == "error" and ws["A1"].font.bold

    def test_template_csv_and_xlsx(self):
        resp = render_template(SPEC)
        text = resp.content.decode("utf-8-sig").split("\r\n")
        assert text[0] == "Nombre,Correo,Asunto,Monto,Fecha,Activo"
        assert text[1] == "Ana Pérez,ana@x.mx,Informes,150.00,05/03/2026,sí"
        ws = load_workbook(io.BytesIO(render_template(SPEC, "xlsx").content)).worksheets[0]
        assert ws["A1"].value == "Nombre" and ws["A2"].value == "Ana Pérez"


# ── the views end-to-end ──────────────────────────────────
class MessagesImportView(ImportView):
    spec = SPEC
    template_url = "/api/v1/x/import/template/"


class MessagesTemplateView(ImportTemplateView):
    spec = SPEC


def _post(data, user=None):
    request = APIRequestFactory().post("/x/import/", data, format="multipart")
    force_authenticate(request, user=user or AdminFactory())
    return MessagesImportView.as_view()(request)


CSV = "Nombre,Correo\nAna,ana@x.mx\n,sin@x.mx\n"


class TestImportView:
    def test_throttle_scope(self):
        assert MessagesImportView.throttle_scope == "admin-import"

    def test_template_endpoint(self):
        request = APIRequestFactory().get("/x/import/template/", {"fmt": "xlsx"})
        force_authenticate(request, user=AdminFactory())
        resp = MessagesTemplateView.as_view()(request)
        assert resp.status_code == 200 and resp["Content-Type"].endswith("spreadsheetml.sheet")
        request = APIRequestFactory().get("/x/import/template/", {"fmt": "pdf"})
        force_authenticate(request, user=AdminFactory())
        assert MessagesTemplateView.as_view()(request).status_code == 400

    def test_dry_run_is_the_default_and_writes_nothing(self):
        resp = _post({"file": _upload(CSV)})
        assert resp.status_code == 200
        assert resp.data["dry_run"] is True and resp.data["entity"] == "mensajes"
        assert resp.data["counts"] == {"crear": 1, "actualizar": 0, "omitir": 0, "error": 1}
        assert resp.data["rows"][0] == {
            "line": 2,
            "key": "ana@x.mx",
            "action": "crear",
            "errors": [],
            "warnings": [],
            "data": {
                "nombre": "Ana",
                "correo": "ana@x.mx",
                "asunto": "",
                "monto": None,
                "fecha": None,
                "activo": None,
            },
        }
        assert ContactMessage.objects.count() == 0
        assert not AuditLog.objects.filter(object_type="import:mensajes").exists()

    def test_report_param_returns_the_annotated_file(self):
        resp = _post({"file": _upload(CSV), "report": "csv"})
        assert resp.status_code == 200 and resp["Content-Type"].startswith("text/csv")
        assert "Falta Nombre." in resp.content.decode("utf-8-sig")
        resp = _post({"file": _upload(CSV), "report": "xlsx"})
        assert resp["Content-Type"].endswith("spreadsheetml.sheet")

    def test_commit_refuses_error_rows_unless_only_valid(self):
        resp = _post({"file": _upload(CSV), "dry_run": "0"})
        assert resp.status_code == 400
        assert "fila(s) con error" in resp.data["detail"]
        assert ContactMessage.objects.count() == 0
        resp = _post({"file": _upload(CSV), "dry_run": "0", "only_valid": "1"})
        assert resp.status_code == 200 and resp.data["dry_run"] is False
        assert ContactMessage.objects.filter(email="ana@x.mx").exists()
        assert AuditLog.objects.filter(object_type="import:mensajes").count() == 1

    def test_missing_file_and_missing_headers_are_400_with_help(self):
        resp = _post({})
        assert resp.status_code == 400
        assert resp.data["detail"].startswith("Adjunte un archivo")
        assert resp.data["template"] == "/api/v1/x/import/template/"
        resp = _post({"file": _upload("Nombre\nAna\n")})
        assert resp.status_code == 400
        assert resp.data["missing"] == ["Correo"]
        assert resp.data["expected_headers"] == SPEC.headers()
        assert resp.data["template"] == "/api/v1/x/import/template/"

    def test_caps_are_413(self):
        class Tiny(MessagesImportView):
            spec = ImportSpec("m", SPEC.columns, max_rows=2, max_bytes=50)

        def post(content):
            request = APIRequestFactory().post(
                "/x/import/", {"file": _upload(content)}, format="multipart"
            )
            force_authenticate(request, user=AdminFactory())
            return Tiny.as_view()(request)

        resp = post("Nombre,Correo\nA,a@x.mx\nB,b@x.mx\nC,c@x.mx\n")
        assert resp.status_code == 413
        assert resp.data == {"detail": "Divida el archivo: el límite es 2 filas."}
        resp = post("Nombre,Correo\n" + "A,a@x.mx\n" * 20)
        assert resp.status_code == 413 and "MB" in resp.data["detail"]
        assert cap_message(2000) == "Acote los filtros: el límite es 2,000 filas."

    def test_non_admin_is_403(self, parent_user):
        assert _post({"file": _upload(CSV)}, user=parent_user).status_code == 403
