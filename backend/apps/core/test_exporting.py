"""
apps.core.exporting — CSV (BOM + CRLF, streamed), XLSX (openpyxl round-trip),
PDF (table_pdf), the 10,000 / 1,000 row caps → 413, ``?ids=`` and the
AdminExportMixin end-to-end (subclassing a list view, audited, throttled).
"""

import io
from datetime import date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from rest_framework import generics
from rest_framework.exceptions import ValidationError
from rest_framework.request import Request
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.accounts.factories import AdminFactory
from apps.core.exceptions import PayloadTooLarge
from apps.core.exporting import (
    AdminExportMixin,
    Col,
    ExportSpec,
    collect_rows,
    parse_ids,
    render_csv,
    render_pdf,
    render_xlsx,
    resolve_path,
    text_value,
    xlsx_value,
)
from apps.core.listing import AdminListMixin
from apps.core.models import AuditLog, ContactMessage
from apps.core.serializers import ContactMessageAdminSerializer

MX = ZoneInfo("America/Mexico_City")

SPEC = ExportSpec(
    filename_prefix="mensajes",
    columns=[
        Col("name", "Nombre", width=20),
        Col("email", "Correo"),
        Col("created_at", "Fecha", fmt="datetime"),
        Col("amount", "Monto", getter=lambda r: Decimal("1234.5"), fmt="money", width=12),
        Col("is_handled", "Atendido"),
    ],
    audit_entity="contact",
    title="Mensajes de contacto",
)


def _rows(n=2):
    return [
        ContactMessage.objects.create(
            name=f"Persona {i}",
            email=f"p{i}@x.mx",
            subject="s",
            message="m",
            is_handled=(i % 2 == 0),
        )
        for i in range(n)
    ]


class TestFormatting:
    def test_resolve_path_walks_relations_and_dicts(self):
        assert resolve_path({"a": {"b": 3}}, "a__b") == 3
        assert resolve_path({"a": None}, "a__b") is None

    def test_text_value_csv_and_human(self):
        col = Col("x", "X")
        assert text_value(col, Decimal("1234.5")) == "1234.50"
        assert text_value(col, Decimal("1234.5"), human=True) == "1,234.50"
        assert text_value(col, True) == "Sí" and text_value(col, False) == "No"
        assert text_value(col, None) == ""
        when = datetime(2026, 3, 5, 14, 7, tzinfo=MX)
        assert text_value(col, when) == "2026-03-05 14:07"
        assert text_value(col, when, human=True) == "05/03/2026 14:07"
        assert text_value(col, date(2026, 3, 5)) == "2026-03-05"
        assert text_value(col, date(2026, 3, 5), human=True) == "05/03/2026"
        assert text_value(Col("x", "X", fmt="int"), 7) == "7"

    def test_xlsx_value_uses_real_types_and_formats(self):
        col = Col("x", "X")
        value, fmt = xlsx_value(col, Decimal("10"))
        assert value == Decimal("10.00") and fmt == "#,##0.00"
        value, fmt = xlsx_value(col, datetime(2026, 3, 5, 14, 7, tzinfo=MX))
        assert value == datetime(2026, 3, 5, 14, 7) and value.tzinfo is None
        assert fmt == "DD/MM/YYYY HH:MM"
        assert xlsx_value(col, date(2026, 3, 5)) == (date(2026, 3, 5), "DD/MM/YYYY")
        assert xlsx_value(col, None) == (None, None)


@pytest.mark.django_db
class TestRenderers:
    def test_csv_streams_with_bom_and_crlf(self):
        rows = _rows(2)
        resp = render_csv(iter(rows), SPEC)
        assert resp["Content-Type"] == "text/csv; charset=utf-8"
        assert resp["Content-Disposition"].startswith('attachment; filename="mensajes_')
        assert resp["Content-Disposition"].endswith('.csv"')
        body = b"".join(resp.streaming_content)
        assert body.startswith("﻿".encode())
        lines = body.decode("utf-8-sig").split("\r\n")
        assert lines[0] == "Nombre,Correo,Fecha,Monto,Atendido"
        assert lines[1].startswith("Persona 0,p0@x.mx,")
        assert lines[1].endswith(",1234.50,Sí")
        assert lines[2].endswith(",1234.50,No")

    def test_xlsx_round_trip(self):
        _rows(2)
        resp = render_xlsx(ContactMessage.objects.order_by("pk"), SPEC)
        assert resp["Content-Type"].endswith("spreadsheetml.sheet")
        assert resp["Content-Disposition"].endswith('.xlsx"')
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.worksheets[0]
        assert ws["A1"].value == "Nombre" and ws["A1"].font.bold
        assert ws.freeze_panes == "A2"
        assert ws.column_dimensions["A"].width == 20
        assert ws["A2"].value == "Persona 0"
        assert isinstance(ws["C2"].value, datetime)
        assert ws["D2"].value == 1234.5 and ws["D2"].number_format == "#,##0.00"
        assert ws["E2"].value == "Sí" and ws["E3"].value == "No"
        assert ws.max_row == 3

    def test_pdf_has_title_header_and_page_numbers(self):
        _rows(3)
        resp = render_pdf(ContactMessage.objects.all(), SPEC)
        assert resp["Content-Type"] == "application/pdf"
        body = resp.content
        assert body.startswith(b"%PDF-1.4")
        assert b"Mensajes de contacto" in body
        assert b"Nombre" in body and b"Persona 0" in body
        assert b"Pagina 1 de 1" in body
        assert b"/MediaBox [0 0 792 612]" in body


@pytest.mark.django_db
class TestCapsAndIds:
    def test_collect_rows_413_above_cap_with_es_mx_message(self):
        _rows(3)
        with pytest.raises(PayloadTooLarge) as exc:
            collect_rows(ContactMessage.objects.all(), 2)
        assert exc.value.status_code == 413
        assert str(exc.value.detail) == "Acote los filtros: el límite es 2 filas."
        count, rows = collect_rows(ContactMessage.objects.all(), 3)
        assert count == 3 and len(list(rows)) == 3

    def test_cap_message_uses_thousands_separator(self):
        assert (
            str(PayloadTooLarge(10_000).detail) == "Acote los filtros: el límite es 10,000 filas."
        )

    def test_parse_ids(self):
        factory = APIRequestFactory()

        def req(ids):
            return Request(factory.get("/x/", {"ids": ids}))

        assert parse_ids(req("3, 1,3,2")) == [3, 1, 2]
        assert parse_ids(Request(factory.get("/x/"))) == []
        with pytest.raises(ValidationError):
            parse_ids(req("1,x"))
        with pytest.raises(PayloadTooLarge):
            parse_ids(req(",".join(str(i) for i in range(1, 502))))


# ── the mixin end-to-end ──────────────────────────────────
class ContactListView(AdminListMixin, generics.ListAPIView):
    serializer_class = ContactMessageAdminSerializer
    queryset = ContactMessage.objects.all()
    search_fields = ("name", "email")
    ordering = {"name": "name"}
    default_ordering = "name"


class ContactExportView(AdminExportMixin, ContactListView):
    export_spec = SPEC


def _export(params, user=None):
    request = APIRequestFactory().get("/x/export/", params)
    force_authenticate(request, user=user or AdminFactory())
    return ContactExportView.as_view()(request)


@pytest.mark.django_db
class TestAdminExportMixin:
    def test_throttle_scope_and_no_pagination(self):
        assert ContactExportView.throttle_scope == "admin-export"
        assert ContactExportView.pagination_class is None

    def test_reuses_list_filters_search_and_ordering(self):
        _rows(3)
        resp = _export({"fmt": "csv", "q": "p1@x.mx", "ordering": "-name"})
        assert resp.status_code == 200
        lines = b"".join(resp.streaming_content).decode("utf-8-sig").strip().split("\r\n")
        assert len(lines) == 2 and lines[1].startswith("Persona 1,")

    def test_ids_exports_only_the_selection(self):
        rows = _rows(3)
        resp = _export({"fmt": "xlsx", "ids": f"{rows[2].pk},{rows[0].pk}"})
        ws = load_workbook(io.BytesIO(resp.content)).worksheets[0]
        assert [ws.cell(row=r, column=1).value for r in (2, 3)] == ["Persona 0", "Persona 2"]

    def test_audits_every_export(self):
        _rows(2)
        _export({"fmt": "pdf", "q": "x.mx"})
        entry = AuditLog.objects.get(object_type="export:contact")
        assert entry.action == "export" and entry.object_id == "pdf"
        assert entry.changes["rows"] == 2 and entry.changes["fmt"] == "pdf"
        assert entry.changes["filters"] == {"q": "x.mx"}
        assert entry.actor_label.startswith("admin")

    def test_bad_format_is_400_and_cap_is_413(self):
        _rows(2)
        assert _export({"fmt": "docx"}).status_code == 400
        small = ExportSpec("m", SPEC.columns, "contact", pdf_row_cap=1)

        class Tiny(ContactExportView):
            export_spec = small

        request = APIRequestFactory().get("/x/export/", {"fmt": "pdf"})
        force_authenticate(request, user=AdminFactory())
        resp = Tiny.as_view()(request)
        assert resp.status_code == 413
        assert resp.data == {"detail": "Acote los filtros: el límite es 1 filas."}
        assert not AuditLog.objects.filter(object_type="export:contact").exists()

    def test_non_admin_cannot_export(self, parent_user):
        assert _export({"fmt": "csv"}, user=parent_user).status_code == 403
