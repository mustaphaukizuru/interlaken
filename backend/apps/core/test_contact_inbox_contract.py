"""
Data Ops Phase 7: the Mensajes inbox on the list contract, its export, the bulk
``mark_handled`` / ``reopen`` actions and the audited single-row toggle.
"""

import pytest
from django.urls import reverse

from apps.accounts.factories import AdminFactory, ParentFactory
from apps.core.filters import CONTACT_ORDERING
from apps.core.models import AuditLog, ContactMessage

pytestmark = pytest.mark.django_db

LIST = reverse("core-contact-inbox")
EXPORT = reverse("core-contact-inbox-export")
BULK = reverse("core-contact-inbox-bulk")

# Mirrors CONTACT_ORDERING_KEYS in frontend/src/hooks/queries/contactMessages.ts.
FRONTEND_ORDERING_KEYS = {"date", "name", "email", "subject", "handled"}


@pytest.fixture
def inbox():
    return [
        ContactMessage.objects.create(
            name="Ana", email="ana@t.mx", subject="Informes", message="Hola"
        ),
        ContactMessage.objects.create(
            name="Beto", email="beto@t.mx", subject="[Facturación] RFC", message="Factura por favor"
        ),
        ContactMessage.objects.create(
            name="Carla", email="carla@t.mx", subject="Visita", message="m", is_handled=True
        ),
    ]


def _ids(resp):
    return [r["id"] for r in resp.data["results"]]


def test_ordering_keys_match_the_frontend_constant():
    assert set(CONTACT_ORDERING) == FRONTEND_ORDERING_KEYS


class TestList:
    def test_parent_is_403(self, api_client):
        api_client.force_authenticate(ParentFactory())
        for url in (LIST, EXPORT):
            assert api_client.get(url).status_code == 403
        assert (
            api_client.post(BULK, {"action": "reopen", "ids": [1]}, format="json").status_code
            == 403
        )

    def test_q_searches_the_message_body_too(self, admin_client, inbox):
        assert _ids(admin_client.get(LIST, {"q": "factura por"})) == [inbox[1].id]
        assert _ids(admin_client.get(LIST, {"q": "carla@"})) == [inbox[2].id]

    def test_handled_filter_and_ordering(self, admin_client, inbox):
        assert set(_ids(admin_client.get(LIST, {"handled": "0"}))) == {inbox[0].id, inbox[1].id}
        assert _ids(admin_client.get(LIST, {"handled": "true"})) == [inbox[2].id]
        assert admin_client.get(LIST, {"handled": "all"}).data["count"] == 3
        assert _ids(admin_client.get(LIST, {"ordering": "-name"})) == [
            inbox[2].id,
            inbox[1].id,
            inbox[0].id,
        ]
        assert _ids(admin_client.get(LIST, {"ordering": "-handled"}))[0] == inbox[2].id

    def test_date_range(self, admin_client, inbox):
        assert admin_client.get(LIST, {"from": "2999-01-01"}).data["count"] == 0

    @pytest.mark.parametrize("n", [3, 12])
    def test_query_budget_is_constant(self, api_client, django_assert_num_queries, n):
        for i in range(n):
            ContactMessage.objects.create(name=f"N{i}", email="x@t.mx", subject="s", message="m")
        api_client.force_authenticate(AdminFactory())
        with django_assert_num_queries(2):
            assert api_client.get(LIST, {"q": "x", "ordering": "name"}).status_code == 200


class TestToggle:
    def test_single_toggle_is_audited(self, admin_client, inbox):
        url = reverse("core-contact-handle", args=[inbox[0].id])
        resp = admin_client.patch(url, {"is_handled": True}, format="json")
        assert resp.status_code == 200 and resp.data["is_handled"] is True
        entry = AuditLog.objects.get(object_type="core.contactmessage", object_id=str(inbox[0].id))
        assert entry.changes == {"is_handled": [False, True]} and entry.context == "contact.handled"

    def test_no_op_toggle_writes_no_audit(self, admin_client, inbox):
        url = reverse("core-contact-handle", args=[inbox[2].id])
        admin_client.patch(url, {"is_handled": True}, format="json")
        assert not AuditLog.objects.filter(object_type="core.contactmessage").exists()


class TestBulk:
    def test_dry_run_plans_and_skips_already_handled(self, admin_client, inbox):
        ids = [m.id for m in inbox]
        resp = admin_client.post(
            BULK, {"action": "mark_handled", "ids": ids, "dry_run": True}, format="json"
        )
        assert resp.status_code == 200
        assert resp.data["ok"] == 2 and [s["id"] for s in resp.data["skipped"]] == [inbox[2].id]
        assert ContactMessage.objects.filter(is_handled=True).count() == 1  # nothing changed

    def test_mark_handled_then_reopen_with_audit(self, admin_client, inbox):
        ids = [m.id for m in inbox]
        resp = admin_client.post(BULK, {"action": "mark_handled", "ids": ids}, format="json")
        assert resp.data["ok"] == 2 and len(resp.data["skipped"]) == 1
        assert ContactMessage.objects.filter(is_handled=True).count() == 3
        per_row = AuditLog.objects.filter(object_type="core.contactmessage")
        assert per_row.count() == 2
        assert AuditLog.objects.filter(
            object_type="bulk:core.contactmessage", object_id="mark_handled"
        ).exists()

        resp = admin_client.post(BULK, {"action": "reopen", "ids": [inbox[0].id]}, format="json")
        assert resp.data["ok"] == 1
        inbox[0].refresh_from_db()
        assert inbox[0].is_handled is False

    def test_all_matching_uses_the_list_filters(self, admin_client, inbox):
        resp = admin_client.post(
            BULK,
            {"action": "mark_handled", "all_matching": True, "filters": {"q": "factura"}},
            format="json",
        )
        assert resp.data["requested"] == 1 and resp.data["ok"] == 1
        inbox[1].refresh_from_db()
        assert inbox[1].is_handled is True
        inbox[0].refresh_from_db()
        assert inbox[0].is_handled is False

    def test_unknown_action_is_400(self, admin_client, inbox):
        resp = admin_client.post(BULK, {"action": "delete", "ids": [inbox[0].id]}, format="json")
        assert resp.status_code == 400


class TestExport:
    def test_csv_filters_and_audit(self, admin_client, inbox):
        resp = admin_client.get(EXPORT, {"fmt": "csv", "handled": "0", "ordering": "name"})
        lines = b"".join(resp.streaming_content).decode("utf-8-sig").strip().split("\r\n")
        assert lines[0] == "Fecha,Nombre,Correo,Asunto,Mensaje,Atendido"
        assert len(lines) == 3 and lines[1].split(",")[1] == "Ana"
        assert AuditLog.objects.get(object_type="export:contact_messages").changes["rows"] == 2

    def test_selected_ids(self, admin_client, inbox):
        resp = admin_client.get(EXPORT, {"fmt": "xlsx", "ids": f"{inbox[2].id}"})
        assert resp.status_code == 200
        from io import BytesIO

        from openpyxl import load_workbook

        ws = load_workbook(BytesIO(resp.content)).active
        assert ws.max_row == 2 and ws["B2"].value == "Carla"
