"""
core/exporting.py — CSV / XLSX / PDF export of an admin list (Data Ops C3).

An export is the list view itself with another renderer: the export view
subclasses the list view (``class AdminFooExportView(AdminExportMixin,
AdminFooListView)``) so filters, search and ordering are shared code, never
duplicated. ``?fmt=csv|xlsx|pdf`` picks the renderer; ``?ids=1,2,3`` (≤ 500)
exports only the selected rows.

Everything runs inside one request (no workers, gunicorn ``--timeout 60``), so
each renderer has a hard row cap: 10,000 for CSV/XLSX, 1,000 for PDF. Above
it the API answers 413 ``{"detail": "Acote los filtros: el límite es N filas."}``.
CSV streams (UTF-8 BOM + CRLF, so Excel es-MX opens it with accents intact);
XLSX uses openpyxl's write-only workbook (bold frozen header, real dates,
money as numbers with ``#,##0.00``); PDF uses ``apps.core.pdf.table_pdf``.

Every export is audited with ``apps.core.audit.record_export``.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from django.http import HttpResponse, StreamingHttpResponse
from django.utils import timezone
from rest_framework.exceptions import ValidationError

from .audit import record_export
from .exceptions import PayloadTooLarge
from .exports import as_download, export_filename, fmt_dt
from .pdf import table_pdf
from .throttling import SharedScopedRateThrottle

ROW_CAP = 10_000
PDF_ROW_CAP = 1_000
IDS_CAP = 500
FORMATS = ("csv", "xlsx", "pdf")
CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
XLSX_DATETIME = "DD/MM/YYYY HH:MM"
XLSX_DATE = "DD/MM/YYYY"
XLSX_MONEY = "#,##0.00"


# ── spec ──────────────────────────────────────────────────
def resolve_path(obj, path: str):
    """``'student__user__email'`` → nested attribute/dict lookup; ``None`` short-circuits."""
    for part in path.split("__"):
        if obj is None:
            return None
        obj = obj.get(part) if isinstance(obj, dict) else getattr(obj, part, None)
        if callable(obj) and not isinstance(obj, type):
            obj = obj()
    return obj


@dataclass(frozen=True)
class Col:
    """One export column. ``fmt``: money | int | date | datetime | bool | text (inferred when None)."""

    key: str
    header: str
    getter: Callable[[Any], Any] | None = None
    width: int | None = None
    fmt: str | None = None

    def value(self, row):
        return self.getter(row) if self.getter else resolve_path(row, self.key)


@dataclass
class ExportSpec:
    filename_prefix: str
    columns: list[Col]
    audit_entity: str
    row_cap: int = ROW_CAP
    pdf_row_cap: int = PDF_ROW_CAP
    title: str = ""
    sheet_title: str = "Datos"
    extra: dict = field(default_factory=dict)

    def cap_for(self, fmt: str) -> int:
        return self.pdf_row_cap if fmt == "pdf" else self.row_cap

    def headers(self) -> list[str]:
        return [c.header for c in self.columns]


# ── value formatting ──────────────────────────────────────
def infer_fmt(col: Col, value) -> str:
    if col.fmt:
        return col.fmt
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, (Decimal, float)):
        return "money"
    if isinstance(value, int):
        return "int"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    return "text"


def text_value(col: Col, value, *, human: bool = False) -> str:
    """CSV (``human=False``: ISO-ish, Excel-parseable) or PDF (``human=True``: DD/MM/YYYY)."""
    if value is None or value == "":
        return ""
    kind = infer_fmt(col, value)
    if kind == "bool":
        return "Sí" if value else "No"
    if kind == "datetime":
        if not isinstance(value, datetime):
            return str(value)
        return timezone.localtime(value).strftime("%d/%m/%Y %H:%M") if human else fmt_dt(value)
    if kind == "date":
        if isinstance(value, datetime):
            value = timezone.localtime(value).date()
        return value.strftime("%d/%m/%Y") if human else value.isoformat()
    if kind == "money":
        amount = Decimal(str(value)).quantize(Decimal("0.01"))
        return f"{amount:,.2f}" if human else f"{amount:.2f}"
    if kind == "int":
        return str(int(value))
    return str(value)


def xlsx_value(col: Col, value) -> tuple[Any, str | None]:
    """``(cell value, number_format)`` with real dates/numbers for Excel."""
    if value is None or value == "":
        return None, None
    kind = infer_fmt(col, value)
    if kind == "bool":
        return ("Sí" if value else "No"), None
    if kind == "datetime" and isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None), XLSX_DATETIME
    if kind == "date":
        if isinstance(value, datetime):
            value = timezone.localtime(value).date()
        return value, XLSX_DATE
    if kind == "money":
        return Decimal(str(value)).quantize(Decimal("0.01")), XLSX_MONEY
    if kind == "int":
        return int(value), "0"
    return str(value), None


# ── renderers ─────────────────────────────────────────────
class _Echo:
    """A write() that returns what it is given: csv.writer → generator chunks."""

    def write(self, value):
        return value


def render_csv(rows: Iterable, spec: ExportSpec) -> StreamingHttpResponse:
    """UTF-8 BOM, CRLF, streamed row by row (``rows`` may be a queryset iterator)."""
    writer = csv.writer(_Echo(), lineterminator="\r\n")

    def stream():
        yield "﻿"
        yield writer.writerow(spec.headers())
        for row in rows:
            yield writer.writerow([text_value(c, c.value(row)) for c in spec.columns])

    response = StreamingHttpResponse(
        (chunk.encode("utf-8") for chunk in stream()), content_type=CONTENT_TYPES["csv"]
    )
    return as_download(response, export_filename(spec.filename_prefix, "csv"))


def xlsx_bytes(headers, rows, *, sheet_title="Datos", widths=None, cells=None) -> bytes:
    """Write-only workbook: bold frozen header, optional widths; ``cells`` maps a
    row to ``[(value, number_format), ...]`` (defaults to plain values)."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet((sheet_title or "Datos")[:31])
    for i, width in enumerate(widths or [], start=1):
        if width:
            ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    bold = Font(bold=True)
    header_cells = []
    for h in headers:
        cell = WriteOnlyCell(ws, value=str(h))
        cell.font = bold
        header_cells.append(cell)
    ws.append(header_cells)
    for row in rows:
        out = []
        for value, number_format in (cells(row) if cells else [(v, None) for v in row]):
            cell = WriteOnlyCell(ws, value=value)
            if number_format:
                cell.number_format = number_format
            out.append(cell)
        ws.append(out)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_xlsx(rows: Iterable, spec: ExportSpec) -> HttpResponse:
    payload = xlsx_bytes(
        spec.headers(),
        rows,
        sheet_title=spec.sheet_title,
        widths=[c.width for c in spec.columns],
        cells=lambda row: [xlsx_value(c, c.value(row)) for c in spec.columns],
    )
    response = HttpResponse(payload, content_type=CONTENT_TYPES["xlsx"])
    return as_download(response, export_filename(spec.filename_prefix, "xlsx"))


def render_pdf(rows: Iterable, spec: ExportSpec, *, subtitle: str = "") -> HttpResponse:
    table = [[text_value(c, c.value(row), human=True) for c in spec.columns] for row in rows]
    pdf = table_pdf(
        spec.title or spec.filename_prefix.replace("_", " ").replace("-", " ").capitalize(),
        spec.headers(),
        table,
        subtitle=subtitle or f"{len(table)} filas",
        widths=[c.width for c in spec.columns],
    )
    response = HttpResponse(pdf, content_type=CONTENT_TYPES["pdf"])
    return as_download(response, export_filename(spec.filename_prefix, "pdf"))


def render_export(rows: Iterable, spec: ExportSpec, fmt: str) -> HttpResponse:
    if fmt == "csv":
        return render_csv(rows, spec)
    if fmt == "xlsx":
        return render_xlsx(rows, spec)
    if fmt == "pdf":
        return render_pdf(rows, spec)
    raise ValidationError({"fmt": ["Use csv, xlsx o pdf."]})


# ── request helpers ───────────────────────────────────────
def parse_fmt(request, allowed=FORMATS) -> str:
    fmt = (request.query_params.get("fmt") or "csv").strip().lower()
    if fmt not in allowed:
        raise ValidationError({"fmt": [f'Use {", ".join(allowed)}.']})
    return fmt


def parse_ids(request, *, cap: int = IDS_CAP) -> list[int]:
    """``?ids=1,2,3`` → ``[1, 2, 3]``; empty when absent; 413 above ``cap``."""
    raw = (request.query_params.get("ids") or "").strip()
    if not raw:
        return []
    ids = []
    for token in raw.replace(" ", "").split(","):
        if not token:
            continue
        if not token.isdigit():
            raise ValidationError({"ids": ["Use una lista de ids separados por coma."]})
        ids.append(int(token))
    if len(ids) > cap:
        raise PayloadTooLarge(cap)
    return list(dict.fromkeys(ids))


def collect_rows(qs, cap: int, *, chunk_size: int = 500):
    """``(count, iterator)`` bounded at ``cap``; 413 when the filtered set is larger."""
    count = qs.count()
    if count > cap:
        raise PayloadTooLarge(cap)
    return count, qs[:cap].iterator(chunk_size=chunk_size)


def export_response(request, qs, spec: ExportSpec, *, formats=FORMATS, context: str = ""):
    """Render ``qs`` (already scoped, filtered and ordered) as ``?fmt=`` and audit it.

    The function form of ``AdminExportMixin.list`` for views that are not on
    the admin list contract, e.g. the family portal exports, which subclass
    their own family-scoped list view and call this from ``list()``. Same
    caps (413 over ``spec.cap_for(fmt)``), same ``?ids=`` narrowing, same
    ``record_export`` row.
    """
    fmt = parse_fmt(request, formats)
    ids = parse_ids(request)
    if ids:
        qs = qs.filter(pk__in=ids)
    count, rows = collect_rows(qs, spec.cap_for(fmt))
    response = render_export(rows, spec, fmt)
    filters = {k: v for k, v in request.query_params.lists() if k != "fmt"}
    record_export(spec.audit_entity, fmt, filters, count, request.user, context=context)
    return response


class AdminExportMixin:
    """Turn a list view into its export sibling.

    ``class AdminFooExportView(AdminExportMixin, AdminFooListView): export_spec = SPEC``
    keeps the list's ``get_queryset``/``filter_queryset`` (filters, ``q``,
    ``ordering``) and answers ``GET ?fmt=&ids=`` with the file. Throttled
    under ``admin-export`` (30/min) and audited via ``record_export``.
    """

    export_spec: ExportSpec | None = None
    export_formats = FORMATS
    throttle_classes = [SharedScopedRateThrottle]
    throttle_scope = "admin-export"
    pagination_class = None

    def get_export_spec(self) -> ExportSpec:
        if self.export_spec is None:
            raise NotImplementedError("export_spec is required")
        return self.export_spec

    def get_export_queryset(self):
        return self.filter_queryset(self.get_queryset())

    def list(self, request, *args, **kwargs):
        fmt = parse_fmt(request, self.export_formats)
        spec = self.get_export_spec()
        qs = self.get_export_queryset()
        ids = parse_ids(request)
        if ids:
            qs = qs.filter(pk__in=ids)
        count, rows = collect_rows(qs, spec.cap_for(fmt))
        response = render_export(rows, spec, fmt)
        filters = {k: v for k, v in request.query_params.lists() if k != "fmt"}
        record_export(spec.audit_entity, fmt, filters, count, request.user)
        return response

    def get(self, request, *args, **kwargs):
        return self.list(request, *args, **kwargs)
