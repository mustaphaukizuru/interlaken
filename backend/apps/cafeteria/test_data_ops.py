"""Data Ops Phase 6: the admin Cafetería lists, exports, bulk actions and the ajuste masivo.

Money invariants under test: no bulk path leaves a wallet negative, the top-up
reference keeps a re-applied recarga idempotent, every row writes an audit
entry and every call writes one summary. Nothing here reaches Loyverse: the
remote reads are patched on ``apps.cafeteria.services`` and the token is blank.
"""

import io
from datetime import timedelta
from decimal import Decimal
from unittest.mock import patch

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.factories import AdminFactory, ParentFactory, StudentProfileFactory
from apps.accounts.models import StudentProfile
from apps.cafeteria import admin_data
from apps.cafeteria.models import (
    BalanceAdjustment,
    CafeteriaBalance,
    CafeteriaTransaction,
    LoyverseProfile,
    TopUpRequest,
)
from apps.core.models import AuditLog
from apps.payments.models import Payment
from apps.portal.models import Notification

pytestmark = pytest.mark.django_db


def _balance(student, amount, threshold="50"):
    return CafeteriaBalance.objects.create(
        student=student, balance=Decimal(str(amount)), low_balance_threshold=Decimal(threshold)
    )


def _student(first="Ana", last="López", **kw):
    s = StudentProfileFactory(**kw)
    s.user.first_name, s.user.last_name = first, last
    s.user.save(update_fields=["first_name", "last_name"])
    return s


@pytest.fixture
def admin(api_client):
    user = AdminFactory()
    api_client.force_authenticate(user)
    return user


def _ids(resp, key="id"):
    return [r[key] for r in resp.json()["results"]]


def _csv(resp) -> str:
    return b"".join(resp.streaming_content).decode("utf-8-sig")


def _upload(text: str, name="ajustes.csv"):
    f = io.BytesIO(text.encode("utf-8"))
    f.name = name
    return f


# ── ordering whitelists = the frontend sortKeys (C2) ─────
def test_ordering_whitelists_match_the_console_sort_keys():
    assert set(admin_data.BALANCE_ORDERING) == {
        "student",
        "grade",
        "balance",
        "last_synced",
        "threshold",
    }
    assert set(admin_data.TRANSACTION_ORDERING) == {"date", "amount", "type", "balance", "student"}
    assert set(admin_data.TOPUP_ORDERING) == {"created_at", "amount", "status", "method", "student"}
    assert set(admin_data.CUSTOMER_ORDERING) == {
        "name",
        "code",
        "kind",
        "points",
        "visits",
        "spent",
        "last_visit",
    }
    assert set(admin_data.RECONCILE_ORDERING) == {"student", "matricula", "grade", "local_balance"}


# ── balances ──────────────────────────────────────────────
class TestBalancesList:
    def test_search_filters_and_ordering(self, api_client, admin):
        parent = ParentFactory(email="mama.rivera@example.mx")
        a = _student("Ana", "Alvarez", student_id="09932", grade="4° Primaria", group="A")
        b = _student("Beto", "Rivera", grade="4° Primaria", group="B", parents=[parent])
        c = _student("Caro", "Zapata", grade="1° Secundaria", loyverse_id="")
        c.apply_status(StudentProfile.Status.WITHDRAWN)
        c.save()
        _balance(a, 10)
        _balance(b, 300)
        _balance(c, 80)
        url = reverse("admin-balances")

        def rows(**p):
            return [r["student"]["id"] for r in api_client.get(url, p).json()["results"]]

        assert rows(q="ci09932") == [a.id]
        assert rows(q="09932") == [a.id]
        assert rows(q="mama.rivera") == [b.id]
        assert rows(grade="4° Primaria", group="b") == [b.id]
        assert rows(status="withdrawn") == [c.id]
        assert rows(low_balance="1") == [a.id]
        assert rows(unlinked="true") == [c.id]
        assert rows() == [a.id, b.id, c.id]  # student name
        assert rows(ordering="-balance") == [b.id, c.id, a.id]
        assert rows(ordering="balance", page_size=2) == [a.id, c.id]

    def test_invalid_status_is_a_400_not_a_500(self, api_client, admin):
        assert api_client.get(reverse("admin-balances"), {"status": "nope"}).status_code == 400

    @pytest.mark.parametrize("n", [3, 9])
    def test_query_budget_is_constant(self, api_client, admin, django_assert_num_queries, n):
        for _ in range(n):
            _balance(StudentProfileFactory(), 20)
        with django_assert_num_queries(2):  # count + page (user + loyverse profile joined)
            resp = api_client.get(reverse("admin-balances"), {"q": "User"})
        assert resp.json()["count"] == n

    def test_admin_only(self, api_client):
        api_client.force_authenticate(ParentFactory())
        assert api_client.get(reverse("admin-balances-export")).status_code == 403
        assert api_client.post(reverse("admin-balances-bulk"), {}, format="json").status_code == 403


class TestBalancesExport:
    def test_csv_honours_filters_and_is_audited(self, api_client, admin):
        a = _student("Ana", "Álvarez", grade="4° Primaria")
        b = _student("Beto", "Rivera", grade="1° Secundaria")
        _balance(a, 10)
        _balance(b, 99)
        resp = api_client.get(
            reverse("admin-balances-export"), {"fmt": "csv", "grade": "4° Primaria"}
        )
        assert resp.status_code == 200
        body = _csv(resp)
        assert "Álvarez" in body and "Rivera" not in body
        log = AuditLog.objects.get(object_type="export:cafeteria.balances")
        assert log.changes["rows"] == 1 and log.changes["filters"]["grade"] == "4° Primaria"

    def test_whole_school_xlsx_and_selected_ids(self, api_client, admin):
        rows = [_balance(StudentProfileFactory(), n) for n in (5, 15, 25)]
        resp = api_client.get(reverse("admin-balances-export"), {"fmt": "xlsx"})
        ws = load_workbook(io.BytesIO(resp.content)).active
        assert ws.max_row == 4 and ws["G2"].number_format == "#,##0.00"
        pdf = api_client.get(
            reverse("admin-balances-export"), {"fmt": "pdf", "ids": f"{rows[0].id},{rows[2].id}"}
        )
        assert pdf["Content-Type"] == "application/pdf"
        assert AuditLog.objects.filter(object_type="export:cafeteria.balances").count() == 2
        assert (
            AuditLog.objects.filter(object_type="export:cafeteria.balances", object_id="pdf")
            .get()
            .changes["rows"]
            == 2
        )


class TestBalancesBulk:
    @patch("apps.cafeteria.services.sync_purchases")
    @patch("apps.cafeteria.services.sync_student_balance", return_value=Decimal("40.00"))
    def test_sync_seeds_linked_rows_and_polls_once(self, seed, poll, api_client, admin):
        linked = _balance(StudentProfileFactory(), 0)
        unlinked = _balance(StudentProfileFactory(loyverse_id=""), 0)
        url = reverse("admin-balances-bulk")
        body = {"action": "sync", "ids": [linked.id, unlinked.id]}

        plan = api_client.post(url, {**body, "dry_run": True}, format="json").json()
        assert plan["ok"] == 1 and plan["skipped"][0]["id"] == unlinked.id
        assert seed.call_count == 0 and poll.call_count == 0

        done = api_client.post(url, body, format="json").json()
        assert done["ok"] == 1 and len(done["skipped"]) == 1
        seed.assert_called_once()
        poll.assert_called_once()
        assert AuditLog.objects.filter(object_type="bulk:cafeteria.cafeteriabalance").count() == 1
        assert AuditLog.objects.filter(
            object_type="cafeteria.cafeteriabalance", object_id=str(linked.id)
        ).exists()

    def test_set_threshold_validates_and_audits(self, api_client, admin):
        a = _balance(StudentProfileFactory(), 10, threshold="50")
        b = _balance(StudentProfileFactory(), 10, threshold="80")
        url = reverse("admin-balances-bulk")
        bad = api_client.post(
            url,
            {"action": "set_threshold", "ids": [a.id], "payload": {"threshold": "-5"}},
            format="json",
        )
        assert bad.status_code == 400
        resp = api_client.post(
            url,
            {"action": "set_threshold", "ids": [a.id, b.id], "payload": {"threshold": "80"}},
            format="json",
        ).json()
        assert resp["ok"] == 1 and resp["skipped"][0]["id"] == b.id
        a.refresh_from_db()
        assert a.low_balance_threshold == Decimal("80.00")

    def test_all_matching_uses_the_list_filters(self, api_client, admin):
        low = _balance(_student(grade="2° Primaria"), 5)
        _balance(_student(grade="3° Primaria"), 5)
        resp = api_client.post(
            reverse("admin-balances-bulk"),
            {
                "action": "set_threshold",
                "all_matching": True,
                "filters": {"grade": "2° Primaria"},
                "payload": {"threshold": "20"},
            },
            format="json",
        ).json()
        assert resp["requested"] == 1 and resp["ok"] == 1
        low.refresh_from_db()
        assert low.low_balance_threshold == Decimal("20.00")


# ── ajuste masivo (import) ────────────────────────────────
class TestAjusteMasivo:
    url = staticmethod(lambda: reverse("admin-balances-import"))

    def _post(self, api_client, text, **data):
        return api_client.post(self.url(), {"file": _upload(text), **data}, format="multipart")

    def test_template_csv_and_xlsx(self, api_client, admin):
        csv = api_client.get(reverse("admin-balances-import-template"), {"fmt": "csv"})
        assert csv.content.decode("utf-8-sig").splitlines()[0] == "matricula,monto,motivo"
        xlsx = api_client.get(reverse("admin-balances-import-template"), {"fmt": "xlsx"})
        assert load_workbook(io.BytesIO(xlsx.content)).active["A1"].value == "matricula"

    def test_dry_run_shows_resulting_balance_and_rejects_negative(self, api_client, admin):
        a = _student(student_id="09932")
        b = _student(student_id="10001")
        _balance(a, 100)
        _balance(b, 20)
        text = (
            "matricula,monto,motivo\n"
            "ci09932,50,Beca\n"
            "10001,-30,Corrección\n"
            "99999,10,No existe\n"
            "09932,5,Repetida\n"
            "10001,0,Cero\n"
        )
        report = self._post(api_client, text).json()
        rows = {r["line"]: r for r in report["rows"]}
        assert rows[2]["action"] == "crear" and rows[2]["data"]["saldo_resultante"] == "150.00"
        assert rows[3]["action"] == "error" and "negativo" in rows[3]["errors"][0]
        assert rows[4]["action"] == "error" and "no existe" in rows[4]["errors"][0]
        assert any("Duplicado de la fila 2" in e for e in rows[5]["errors"])
        assert rows[6]["action"] == "error"
        assert report["counts"]["crear"] == 1
        assert BalanceAdjustment.objects.count() == 0  # a dry run writes nothing

    def test_commit_refused_with_errors_unless_valid_only(self, api_client, admin):
        parent = ParentFactory()
        a = _student(student_id="09932", parents=[parent])
        b = _student(student_id="10001")
        _balance(a, 100)
        _balance(b, 20)
        text = "matricula,monto,motivo\n09932,-40,Ajuste\n10001,-30,Excede\n"
        refused = self._post(api_client, text, dry_run="0")
        assert refused.status_code == 400
        with patch("apps.cafeteria.services._session") as remote:
            done = self._post(
                api_client, text, dry_run="0", valid_only="1", notify="0", mirror="0"
            ).json()
        remote.assert_not_called()  # mirror=0: no Loyverse push
        assert done["counts"] == {"crear": 1, "actualizar": 0, "omitir": 0, "error": 1}
        assert CafeteriaBalance.objects.get(student=a).balance == Decimal("60.00")
        assert CafeteriaBalance.objects.get(student=b).balance == Decimal("20.00")
        adj = BalanceAdjustment.objects.get()
        assert adj.amount == Decimal("-40.00") and adj.transaction is not None
        assert adj.transaction.loyverse_receipt_id.startswith("adjust-tx-")
        assert not Notification.objects.filter(user=parent).exists()  # notify=0
        assert AuditLog.objects.filter(
            object_type="import:ajustes_cafeteria", action="import"
        ).exists()
        assert AuditLog.objects.filter(
            object_type="cafeteria.balanceadjustment", context="import:ajustes_cafeteria"
        ).exists()

    def test_notify_default_tells_the_family_and_flags_a_reupload(self, api_client, admin):
        parent = ParentFactory()
        a = _student(student_id="09932", parents=[parent])
        _balance(a, 0)
        text = "matricula,monto,motivo\n09932,25,Beca octubre\n"
        self._post(api_client, text, dry_run="0")
        assert Notification.objects.filter(user=parent).count() == 1
        again = self._post(api_client, text).json()
        assert "últimas 24 h" in again["rows"][0]["warnings"][0]

    def test_error_report_and_row_cap(self, api_client, admin):
        report = self._post(api_client, "matricula,monto,motivo\n12345,1,x\n", report="csv")
        body = report.content.decode("utf-8-sig")
        assert "resultado" in body.splitlines()[0] and "error" in body
        big = "matricula,monto,motivo\n" + "".join(f"{i},1,x\n" for i in range(501))
        assert self._post(api_client, big).status_code == 413


# ── transactions + adjustments ────────────────────────────
class TestTransactions:
    def _tx(self, student, kind, amount, ref, desc="", days=0):
        return CafeteriaTransaction.objects.create(
            student=student,
            transaction_type=kind,
            amount=Decimal(amount),
            description=desc,
            loyverse_receipt_id=ref,
            date=timezone.now() - timedelta(days=days),
        )

    def test_search_filters_and_export(self, api_client, admin):
        a = _student("Ana", "Álvarez", student_id="09932")
        b = _student("Beto", "Rivera")
        t1 = self._tx(a, "purchase", "30", "R-1001", "Torta", days=10)
        t2 = self._tx(b, "topup", "100", "topup-payment-9", "Recarga", days=1)
        url = reverse("admin-transactions")
        assert _ids(api_client.get(url, {"q": "ci09932"})) == [t1.id]
        assert _ids(api_client.get(url, {"q": "R-1001"})) == [t1.id]
        assert _ids(api_client.get(url, {"q": "torta"})) == [t1.id]
        assert _ids(api_client.get(url, {"type": "topup"})) == [t2.id]
        assert _ids(api_client.get(url, {"student": a.id})) == [t1.id]
        since = (timezone.localdate() - timedelta(days=3)).isoformat()
        assert _ids(api_client.get(url, {"from": since})) == [t2.id]
        assert _ids(api_client.get(url, {"ordering": "amount"})) == [t1.id, t2.id]
        row = api_client.get(url).json()["results"][0]
        assert row["student_name"] and row["type_display"]
        body = _csv(
            api_client.get(reverse("admin-transactions-export"), {"fmt": "csv", "type": "purchase"})
        )
        assert "R-1001" in body and "topup-payment-9" not in body

    @pytest.mark.parametrize("n", [2, 6])
    def test_query_budget(self, api_client, admin, django_assert_num_queries, n):
        for i in range(n):
            self._tx(StudentProfileFactory(), "purchase", "5", f"qb-{n}-{i}")
        with django_assert_num_queries(2):
            assert api_client.get(reverse("admin-transactions")).json()["count"] == n

    def test_adjustments_are_paged_per_student(self, api_client, admin):
        a = StudentProfileFactory()
        _balance(a, 100)
        from apps.cafeteria.services import adjust_balance

        for i in range(3):
            adjust_balance(a, Decimal("1"), f"r{i}", notify=False, mirror=False)
        resp = api_client.get(reverse("admin-adjustments"), {"student": a.id, "page_size": 2})
        assert resp.json()["count"] == 3 and len(resp.json()["results"]) == 2

    def test_student_detail_can_skip_the_ledger(self, api_client, admin):
        a = StudentProfileFactory()
        self._tx(a, "purchase", "5", "d-1")
        data = api_client.get(reverse("admin-student-detail", args=[a.id]), {"ledger": "0"}).json()
        assert data["transactions"] == [] and data["adjustments"] == []


# ── top-ups ───────────────────────────────────────────────
class TestTopUps:
    def test_search_ordering_and_export(self, api_client, admin):
        a = _student("Ana", "Álvarez")
        t1 = TopUpRequest.objects.create(student=a, amount=Decimal("100"), payment_ref="CAJA-77")
        t2 = TopUpRequest.objects.create(student=a, amount=Decimal("50"))
        url = reverse("admin-topups")
        assert _ids(api_client.get(url, {"q": "caja-77"})) == [t1.id]
        assert _ids(api_client.get(url, {"ordering": "amount"})) == [t2.id, t1.id]
        body = _csv(api_client.get(reverse("admin-topups-export"), {"fmt": "csv"}))
        assert "CAJA-77" in body

    def test_bulk_apply_credits_office_pending_once(self, api_client, admin):
        student = StudentProfileFactory()
        _balance(student, 10)
        office = TopUpRequest.objects.create(student=student, amount=Decimal("100"))
        online = TopUpRequest.objects.create(
            student=student, amount=Decimal("70"), method=TopUpRequest.Method.ONLINE
        )
        no_link = TopUpRequest.objects.create(
            student=StudentProfileFactory(loyverse_id=""), amount=Decimal("30")
        )
        url = reverse("admin-topups-bulk")
        body = {"action": "apply", "ids": [office.id, online.id, no_link.id]}
        plan = api_client.post(url, {**body, "dry_run": True}, format="json").json()
        assert plan["ok"] == 1 and plan["total"] == "100.00"
        assert [s["id"] for s in plan["skipped"]] == [online.id]
        assert [f["id"] for f in plan["failed"]] == [no_link.id]

        done = api_client.post(url, body, format="json").json()
        assert done["ok"] == 1 and done["total"] == "100.00"
        assert CafeteriaBalance.objects.get(student=student).balance == Decimal("110.00")
        tx = CafeteriaTransaction.objects.get(loyverse_receipt_id=f"topup-request-{office.id}")
        assert tx.transaction_type == CafeteriaTransaction.TxType.TOPUP

        # A second apply is skipped (already completed): no double credit.
        again = api_client.post(url, {"action": "apply", "ids": [office.id]}, format="json").json()
        assert again["ok"] == 0 and len(again["skipped"]) == 1
        assert CafeteriaBalance.objects.get(student=student).balance == Decimal("110.00")
        assert AuditLog.objects.filter(
            object_type="cafeteria.topuprequest", object_id=str(office.id)
        ).exists()
        assert AuditLog.objects.filter(object_type="bulk:cafeteria.topuprequest").count() == 2

    def test_bulk_pos_loaded_and_unloaded(self, api_client, admin):
        student = StudentProfileFactory()
        ready = TopUpRequest.objects.create(
            student=student,
            amount=Decimal("40"),
            method=TopUpRequest.Method.ONLINE,
            status=TopUpRequest.Status.COMPLETED,
        )
        refunded = TopUpRequest.objects.create(
            student=student,
            amount=Decimal("40"),
            method=TopUpRequest.Method.ONLINE,
            status=TopUpRequest.Status.COMPLETED,
        )
        Payment.objects.create(
            user=ParentFactory(),
            payment_type=Payment.Type.CAFETERIA,
            amount=Decimal("40"),
            related_topup=refunded,
            status=Payment.Status.REFUNDED,
        )
        url = reverse("admin-topups-bulk")
        res = api_client.post(
            url, {"action": "pos_loaded", "ids": [ready.id, refunded.id]}, format="json"
        ).json()
        assert res["ok"] == 1 and res["skipped"][0]["id"] == refunded.id
        ready.refresh_from_db()
        assert ready.pos_loaded_by == admin

        ready.pos_unload_needed_at = timezone.now()
        ready.save(update_fields=["pos_unload_needed_at"])
        res = api_client.post(
            url, {"action": "pos_unloaded", "ids": [ready.id, refunded.id]}, format="json"
        ).json()
        assert res["ok"] == 1
        ready.refresh_from_db()
        assert ready.pos_unloaded_at is not None

    def test_bulk_cap_is_413(self, api_client, admin):
        resp = api_client.post(
            reverse("admin-topups-bulk"),
            {"action": "apply", "ids": list(range(1, 502))},
            format="json",
        )
        assert resp.status_code == 413


# ── low balance, customers ────────────────────────────────
class TestLowBalanceAndCustomers:
    def test_low_balance_search_and_export(self, api_client, admin):
        a = _student("Ana", "Álvarez")
        b = _student("Beto", "Rivera")
        for s in (a, b):
            _balance(s, 5)
            CafeteriaTransaction.objects.create(
                student=s,
                transaction_type="purchase",
                amount=Decimal("5"),
                loyverse_receipt_id=f"lb-{s.pk}",
            )
        url = reverse("admin-low-balance")
        assert [
            r["student"]["id"] for r in api_client.get(url, {"q": "rivera"}).json()["results"]
        ] == [b.id]
        body = _csv(api_client.get(reverse("admin-low-balance-export"), {"fmt": "csv"}))
        assert "Álvarez" in body and "Rivera" in body

    def test_customers_page_size_summary_and_export(self, api_client, admin):
        for i in range(3):
            LoyverseProfile.objects.create(loyverse_id=f"c{i}", name=f"Cliente {i}", kind="staff")
        LoyverseProfile.objects.create(
            loyverse_id="gone", name="Borrado", kind="test", missing_since=timezone.now()
        )
        url = reverse("admin-customers")
        data = api_client.get(url, {"page_size": 2, "ordering": "name"}).json()
        assert data["count"] == 4 and len(data["results"]) == 2
        assert data["summary"]["staff"] == 3 and data["summary"]["missing"] == 1
        assert _ids(api_client.get(url, {"missing": "1"}), "loyverse_id") == ["gone"]
        body = _csv(
            api_client.get(reverse("admin-customers-export"), {"fmt": "csv", "kind": "staff"})
        )
        assert "Cliente 0" in body and "Borrado" not in body


# ── reconciliation ────────────────────────────────────────
class TestReconcile:
    @patch("apps.cafeteria.services.get_customer_by_id")
    def test_paged_list_with_search(self, remote, api_client, admin):
        remote.side_effect = lambda lid: {"total_points": 50}
        a = _student("Ana", "Alvarez")
        b = _student("Beto", "Rivera")
        _balance(a, 50)
        _balance(b, 70)
        data = api_client.get(reverse("admin-reconcile"), {"page_size": 1}).json()
        assert data["total"] == 2 and data["checked"] == 1 and data["has_more"] is True
        assert data["results"][0]["student_id"] == a.id
        found = api_client.get(reverse("admin-reconcile"), {"q": "rivera"}).json()
        assert found["results"][0]["drift"] == "20.00"

    @patch("apps.cafeteria.services.get_customer_by_id")
    def test_bulk_fix_sets_local_to_loyverse_without_notifying(self, remote, api_client, admin):
        parent = ParentFactory()
        drift = StudentProfileFactory(parents=[parent])
        ok = StudentProfileFactory()
        negative = StudentProfileFactory()
        _balance(drift, 80)
        _balance(ok, 50)
        _balance(negative, 10)
        points = {drift.loyverse_id: 65, ok.loyverse_id: 50, negative.loyverse_id: -5}
        remote.side_effect = lambda lid: {"total_points": points[lid]}
        url = reverse("admin-reconcile-bulk")
        body = {"action": "fix", "ids": [drift.id, ok.id, negative.id]}
        plan = api_client.post(url, {**body, "dry_run": True}, format="json").json()
        assert plan["ok"] == 1 and len(plan["skipped"]) == 1 and len(plan["failed"]) == 1
        assert CafeteriaBalance.objects.get(student=drift).balance == Decimal("80.00")

        done = api_client.post(url, body, format="json").json()
        assert done["ok"] == 1
        assert CafeteriaBalance.objects.get(student=drift).balance == Decimal("65.00")
        assert CafeteriaBalance.objects.get(student=negative).balance == Decimal("10.00")
        assert not Notification.objects.filter(user=parent).exists()
        assert BalanceAdjustment.objects.get(student=drift).amount == Decimal("-15.00")

    def test_bulk_fix_caps_and_refuses_all_matching(self, api_client, admin):
        url = reverse("admin-reconcile-bulk")
        too_many = api_client.post(url, {"action": "fix", "ids": list(range(1, 52))}, format="json")
        assert too_many.status_code == 413
        all_matching = api_client.post(url, {"action": "fix", "all_matching": True}, format="json")
        assert all_matching.status_code == 400

    @patch("apps.cafeteria.services.get_all_customers")
    def test_export_uses_one_store_fetch(self, store, api_client, admin):
        a = _student("Ana", "Álvarez")
        _balance(a, 30)
        store.return_value = [{"id": a.loyverse_id, "total_points": 25}]
        body = _csv(api_client.get(reverse("admin-reconcile-export"), {"fmt": "csv"}))
        assert "Álvarez" in body and "5.00" in body
        store.assert_called_once()
        assert AuditLog.objects.filter(object_type="export:cafeteria.reconcile").exists()
