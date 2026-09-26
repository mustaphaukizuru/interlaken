"""Data Ops Phase 9: the family cafetería history and its export.

``GET /cafeteria/transactions/`` (list) and ``GET /cafeteria/export/`` (file)
share one queryset: the export subclasses the list view. These tests pin the
contract the portal page relies on: strict family scoping (two families), the
four whitelisted sort keys with the ``-pk`` tiebreak, the ``student``/``type``/
``from``/``to`` filters on both endpoints, ``page_size`` capped at 100, the
10,000-row cap (413), the ``record_export`` audit row and O(1) query counts.
"""

import io
from datetime import datetime
from decimal import Decimal
from itertools import count

import pytest
from django.urls import reverse

from apps.accounts.factories import AdminFactory, ParentFactory, StudentProfileFactory, UserFactory
from apps.accounts.models import User
from apps.cafeteria import views as cafeteria_views
from apps.cafeteria.models import CafeteriaTransaction
from apps.cafeteria.portal_exports import FAMILY_TRANSACTIONS_EXPORT_SPEC
from apps.core.listing import MEXICO_CITY
from apps.core.models import AuditLog

pytestmark = pytest.mark.django_db

_receipts = count(1)

PURCHASE = CafeteriaTransaction.TxType.PURCHASE
TOPUP = CafeteriaTransaction.TxType.TOPUP
REFUND = CafeteriaTransaction.TxType.REFUND


def _tx(student, tx_type, amount, balance, day, description=""):
    return CafeteriaTransaction.objects.create(
        student=student,
        transaction_type=tx_type,
        amount=Decimal(amount),
        balance_after=Decimal(balance),
        description=description,
        loyverse_receipt_id=f"R-phase9-{next(_receipts)}",
        date=datetime(2026, 9, day, 12, 0, tzinfo=MEXICO_CITY),
    )


def _csv(resp) -> str:
    return b"".join(resp.streaming_content).decode("utf-8-sig")


def _csv_rows(resp) -> list[str]:
    return [line for line in _csv(resp).split("\r\n")[1:] if line]


@pytest.fixture
def families():
    """Two families; family A has two children, family B one."""
    parent_a = ParentFactory()
    parent_b = ParentFactory()
    a1 = StudentProfileFactory(parents=[parent_a])
    a2 = StudentProfileFactory(parents=[parent_a])
    b1 = StudentProfileFactory(parents=[parent_b])
    _tx(a1, TOPUP, "200.00", "200.00", 1, "Recarga A1")
    _tx(a1, PURCHASE, "35.50", "164.50", 3, "Torta A1")
    _tx(a1, REFUND, "10.00", "174.50", 5, "Devolucion A1")
    _tx(a2, PURCHASE, "80.00", "20.00", 4, "Comida A2")
    _tx(b1, PURCHASE, "999.00", "1.00", 2, "Secreto B1")
    return {"parent_a": parent_a, "parent_b": parent_b, "a1": a1, "a2": a2, "b1": b1}


class TestFamilyScoping:
    def test_list_shows_only_own_children(self, api_client, families):
        api_client.force_authenticate(families["parent_a"])
        rows = api_client.get(reverse("cafeteria-transactions")).data["results"]
        assert {r["student_id"] for r in rows} == {families["a1"].id, families["a2"].id}
        assert all(r["description"] != "Secreto B1" for r in rows)

    def test_other_familys_student_filter_is_empty_not_leaked(self, api_client, families):
        api_client.force_authenticate(families["parent_a"])
        resp = api_client.get(reverse("cafeteria-transactions"), {"student": families["b1"].id})
        assert resp.status_code == 200 and resp.data["count"] == 0
        export = api_client.get(reverse("cafeteria-export"), {"student": families["b1"].id})
        assert export.status_code == 200 and _csv_rows(export) == []

    def test_export_is_family_scoped_both_ways(self, api_client, families):
        api_client.force_authenticate(families["parent_a"])
        body = _csv(api_client.get(reverse("cafeteria-export")))
        assert "Torta A1" in body and "Comida A2" in body and "Secreto B1" not in body
        assert families["b1"].student_id not in body

        api_client.force_authenticate(families["parent_b"])
        body = _csv(api_client.get(reverse("cafeteria-export")))
        assert "Secreto B1" in body and "Torta A1" not in body and "Comida A2" not in body

    def test_student_login_sees_only_own_profile(self, api_client, families):
        api_client.force_authenticate(families["a1"].user)
        rows = _csv_rows(api_client.get(reverse("cafeteria-export")))
        assert len(rows) == 3 and all(families["a1"].student_id in r for r in rows)

    def test_staff_and_anonymous_are_refused(self, api_client, families):
        assert api_client.get(reverse("cafeteria-export")).status_code == 401
        api_client.force_authenticate(UserFactory(role=User.Role.STAFF))
        assert api_client.get(reverse("cafeteria-export")).status_code == 403


class TestOrdering:
    @pytest.mark.parametrize(
        "ordering, expected",
        [
            ("", ["Devolucion A1", "Comida A2", "Torta A1", "Recarga A1"]),
            ("date", ["Recarga A1", "Torta A1", "Comida A2", "Devolucion A1"]),
            ("-amount", ["Recarga A1", "Comida A2", "Torta A1", "Devolucion A1"]),
            ("amount", ["Devolucion A1", "Torta A1", "Comida A2", "Recarga A1"]),
            ("balance", ["Comida A2", "Torta A1", "Devolucion A1", "Recarga A1"]),
            ("-balance", ["Recarga A1", "Devolucion A1", "Torta A1", "Comida A2"]),
        ],
    )
    def test_list_and_export_follow_the_same_order(self, api_client, families, ordering, expected):
        api_client.force_authenticate(families["parent_a"])
        params = {"ordering": ordering} if ordering else {}
        rows = api_client.get(reverse("cafeteria-transactions"), params).data["results"]
        assert [r["description"] for r in rows] == expected
        exported = _csv_rows(api_client.get(reverse("cafeteria-export"), params))
        assert [next(d for d in expected if d in line) for line in exported] == expected

    def test_type_sort_groups_types(self, api_client, families):
        api_client.force_authenticate(families["parent_a"])
        rows = api_client.get(reverse("cafeteria-transactions"), {"ordering": "type"}).data[
            "results"
        ]
        assert [r["transaction_type"] for r in rows] == ["purchase", "purchase", "refund", "topup"]

    def test_unknown_or_relation_key_falls_back_to_default(self, api_client, families):
        api_client.force_authenticate(families["parent_a"])
        rows = api_client.get(
            reverse("cafeteria-transactions"), {"ordering": "student__user__password"}
        ).data["results"]
        assert rows[0]["description"] == "Devolucion A1"

    def test_whitelist_is_the_four_portal_keys(self):
        # frontend/src/pages/parent/CafeteriaPage.tsx TX_SORT_OPTIONS uses the same four.
        assert set(cafeteria_views.TRANSACTION_ORDERING) == {"date", "amount", "type", "balance"}


class TestFilters:
    def test_type_and_dates_on_list_and_export(self, api_client, families):
        api_client.force_authenticate(families["parent_a"])
        params = {"type": "purchase", "from": "2026-09-03", "to": "2026-09-04"}
        rows = api_client.get(reverse("cafeteria-transactions"), params).data["results"]
        assert sorted(r["description"] for r in rows) == ["Comida A2", "Torta A1"]
        exported = _csv_rows(api_client.get(reverse("cafeteria-export"), params))
        assert len(exported) == 2
        assert all("Torta A1" in r or "Comida A2" in r for r in exported)

    def test_student_filter_on_export(self, api_client, families):
        api_client.force_authenticate(families["parent_a"])
        exported = _csv_rows(
            api_client.get(reverse("cafeteria-export"), {"student": families["a2"].id})
        )
        assert len(exported) == 1 and "Comida A2" in exported[0]

    def test_page_size_is_capped_at_100(self, api_client):
        parent = ParentFactory()
        child = StudentProfileFactory(parents=[parent])
        CafeteriaTransaction.objects.bulk_create(
            [
                CafeteriaTransaction(
                    student=child,
                    transaction_type=PURCHASE,
                    amount=Decimal("1.00"),
                    loyverse_receipt_id=f"R-phase9-cap-{i}",
                )
                for i in range(105)
            ]
        )
        api_client.force_authenticate(parent)
        data = api_client.get(reverse("cafeteria-transactions"), {"page_size": 500}).data
        assert data["count"] == 105 and len(data["results"]) == 100


class TestExportContract:
    def test_csv_shape(self, api_client, families):
        api_client.force_authenticate(families["parent_a"])
        resp = api_client.get(reverse("cafeteria-export"))
        assert resp["Content-Type"].startswith("text/csv")
        assert 'filename="movimientos_cafeteria_' in resp["Content-Disposition"]
        raw = b"".join(resp.streaming_content)
        assert raw.startswith("﻿".encode())
        header = raw.decode("utf-8-sig").split("\r\n")[0]
        assert header == "Alumno,Matrícula,Fecha,Tipo,Descripción,Monto,Saldo"

    def test_xlsx(self, api_client, families):
        from openpyxl import load_workbook

        api_client.force_authenticate(families["parent_a"])
        resp = api_client.get(reverse("cafeteria-export"), {"fmt": "xlsx", "ordering": "date"})
        assert resp.status_code == 200
        ws = load_workbook(io.BytesIO(resp.content)).active
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][0] == "Alumno" and len(rows) == 5
        assert rows[1][4] == "Recarga A1"
        assert float(rows[1][5]) == 200.0  # money is a real number in Excel, not text

    def test_pdf_is_not_offered(self, api_client, families):
        api_client.force_authenticate(families["parent_a"])
        assert api_client.get(reverse("cafeteria-export"), {"fmt": "pdf"}).status_code == 400

    def test_export_is_audited_with_filters_and_count(self, api_client, families):
        parent = families["parent_a"]
        api_client.force_authenticate(parent)
        api_client.get(reverse("cafeteria-export"), {"type": "purchase", "ordering": "-amount"})
        entry = AuditLog.objects.get(object_type="export:cafeteria.family_transactions")
        assert entry.actor == parent and entry.action == "export" and entry.context == "portal"
        assert entry.changes["rows"] == 2 and entry.changes["fmt"] == "csv"
        assert entry.changes["filters"] == {"type": "purchase", "ordering": "-amount"}

    def test_over_cap_is_413(self, api_client, families, monkeypatch):
        monkeypatch.setattr(FAMILY_TRANSACTIONS_EXPORT_SPEC, "row_cap", 2)
        api_client.force_authenticate(families["parent_a"])
        resp = api_client.get(reverse("cafeteria-export"))
        assert resp.status_code == 413
        assert resp.data["detail"] == "Acote los filtros: el límite es 2 filas."
        assert not AuditLog.objects.filter(object_type__startswith="export:").exists()

    def test_throttled_under_portal_scope(self):
        assert cafeteria_views.ParentExportView.throttle_scope == "portal-export"

    def test_admin_keeps_access(self, api_client, families):
        api_client.force_authenticate(AdminFactory())
        resp = api_client.get(reverse("cafeteria-export"), {"student": families["b1"].id})
        assert resp.status_code == 200 and len(_csv_rows(resp)) == 1


# list: family student ids subquery is inlined → count + page rows.
HISTORY_LIST_BUDGET = 2
# export: count + rows (student + user joined) + audit insert.
HISTORY_EXPORT_BUDGET = 3


class TestQueryBudgets:
    @pytest.mark.parametrize("n_rows", [3, 9])
    def test_list_and_export_are_constant(self, api_client, django_assert_num_queries, n_rows):
        parent = ParentFactory()
        kids = [StudentProfileFactory(parents=[parent]) for _ in range(3)]
        for i in range(n_rows):
            _tx(kids[i % 3], PURCHASE, "5.00", "100.00", 1 + i % 20)
        api_client.force_authenticate(parent)
        with django_assert_num_queries(HISTORY_LIST_BUDGET):
            assert api_client.get(reverse("cafeteria-transactions")).data["count"] == n_rows
        with django_assert_num_queries(HISTORY_EXPORT_BUDGET):
            resp = api_client.get(reverse("cafeteria-export"))
            assert len(_csv_rows(resp)) == n_rows
